#!/usr/bin/env python3
"""PySAS 0.3.16 — portable SAS Enterprise Guide command-line utilities.

Keep this file beside the EGP, scheduler workbook and any top-level _*.sas
initialisation files it should use.  Python 3.9+ is recommended.  ``rich`` is
optional; ``openpyxl`` is required only for scheduler and table-workbook work.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import contextlib
import math
import csv
import hashlib
import json
import os
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Any, Iterable
from xml.etree import ElementTree as ET


VERSION = "0.3.16"
ROOT_DIR = Path(__file__).resolve().parent
CODEBASE_FILE = "codebase.sasbundle.txt"
BACKUP_FOLDER = "_codebase_backups"
FORMAT_VERSION = 1
BUNDLE_START = "/*@@SAS_BUNDLE_START@@"
BUNDLE_END = "@@SAS_BUNDLE_END@@*/"
FILE_START = "/*@@SAS_FILE_START@@"
FILE_END = "@@SAS_FILE_END@@*/"
FOOTER_START = "/*@@SAS_BUNDLE_FOOTER@@"
FOOTER_END = "@@SAS_BUNDLE_FOOTER_END@@*/"
REQUIRED_COLUMNS = {
    "task_id", "program", "depends_on", "skip", "row_start", "row_end",
    "section", "stop_process_on_error", "stop_program_on_error",
    "max_parallel", "always_run",
}
FINAL_OK = {"SUCCESS", "SKIPPED_SUCCESS", "SKIPPED_PREVIOUS", "ALWAYS_RUN_DEFINITION"}
CONNECTION_PATTERNS = (
    "the connection to the server has been lost",
    "unable to communicate with the server",
    "server is disconnected",
    "rpc_e_disconnected",
    "remote object has been lost",
    "wrapperlostexception",
    "object invoked has disconnected",
)
PRINT_LOCK = threading.Lock()


def now_stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def text_encoding(raw: bytes) -> str:
    """Only treat text as UTF-16 when its BOM or NUL layout supports it."""
    if raw.startswith(b"\x00MSMAMARPCRYPT"):
        raise ValueError("This file is encrypted/protected (MSMAMARPCRYPT), not readable source text. Use a readable copy from an application authorized to open it.")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    sample = raw[:4096]
    if len(sample) >= 4:
        even, odd = sample[::2], sample[1::2]
        if odd.count(0) / len(odd) > .3 and even.count(0) / len(even) < .1:
            return "utf-16-le"
        if even.count(0) / len(even) > .3 and odd.count(0) / len(odd) < .1:
            return "utf-16-be"
    import codecs
    try:
        codecs.getincrementaldecoder("utf-8-sig")().decode(raw, final=False)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1252"


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    return raw.decode(text_encoding(raw), errors="replace")


def normalized(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").rstrip("\n") + "\n"


def safe_name(value: str, limit: int = 70) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._-") or "item"
    return value[:limit]


def truthy(value: Any) -> bool:
    return str(value or "").strip().casefold() in {"1", "1.0", "true", "yes", "y", "x"}


def require_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        return Workbook, load_workbook, Font, PatternFill, get_column_letter
    except ImportError as exc:
        raise RuntimeError("This command needs openpyxl: py -m pip install openpyxl") from exc


def find_one(pattern: str, description: str, folder: Path = ROOT_DIR) -> Path:
    items = sorted(p for p in folder.glob(pattern) if p.is_file() and not p.name.startswith("~$"))
    if not items:
        raise FileNotFoundError(f"No {description} found beside {Path(__file__).name}")
    if len(items) != 1:
        raise RuntimeError(f"Multiple {description}s found: " + ", ".join(p.name for p in items))
    return items[0]


def sas_files(folder: Path) -> list[Path]:
    """Find SAS files without depending on the filesystem's case sensitivity."""
    return sorted((p for p in folder.iterdir() if p.is_file() and p.suffix.casefold() == ".sas"),
                  key=lambda p: (p.name.casefold(), p.name)) if folder.is_dir() else []


def init_files(explicit: Path | None = None, folder: Path | None = None, extra_folder: Path | None = None) -> list[Path]:
    if explicit:
        if not explicit.is_file():
            raise FileNotFoundError(explicit)
        return [explicit]
    folders = [folder or ROOT_DIR]
    if extra_folder is not None:
        folders.append(extra_folder)
    paths = {p.resolve() for directory in folders for p in sas_files(directory) if p.name.startswith("_")}
    return sorted(paths, key=lambda p: (p.name.casefold(), str(p).casefold()))


def print_home() -> None:
    inits = init_files()
    egps = sorted(ROOT_DIR.glob("*.egp"))
    excels = sorted(p for p in ROOT_DIR.glob("*.xlsx") if not p.name.startswith("~$"))
    print(f"PySAS {VERSION}")
    print("SAS Enterprise Guide command-line utilities\n")
    print(f"Folder: {ROOT_DIR}")
    print(f"EGP projects: {len(egps)} | Excel workbooks: {len(excels)} | SAS files: {len(list(ROOT_DIR.glob('*.sas')))}")
    names = ", ".join(p.name for p in inits) if inits else "none"
    print(f"Runner init: {len(inits)} _*.sas file(s) ({names})")
    print("\nRunner convention:")
    print("  All _*.sas files beside pysas.py run first, alphabetically.")
    print("  A job named foo.tables.sas automatically exports all output datasets.")
    print("\nCommon commands:")
    print("  py pysas.py runner watch")
    print("  py pysas.py runner run diagnostic.sas")
    print("  py pysas.py schedule")
    print("  py pysas.py egp inspect")
    print("  py pysas.py bundle pack")
    print("  py pysas.py --help")


# ---------------------------------------------------------------------------
# Bundle engine
# ---------------------------------------------------------------------------


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def marker(prefix: str, data: dict[str, Any], suffix: str) -> str:
    return f"{prefix} {json.dumps(data, ensure_ascii=False, sort_keys=True)} {suffix}\n"


def discover_sas(recursive: bool, root: Path = ROOT_DIR) -> list[Path]:
    paths = root.glob("**/*.sas" if recursive else "*.sas")
    backup = (root / BACKUP_FOLDER).resolve()
    result = []
    for path in paths:
        if not path.is_file() or not path.resolve().is_relative_to(root.resolve()):
            continue
        try:
            path.resolve().relative_to(backup)
            continue
        except ValueError:
            result.append(path)
    return sorted(result, key=lambda p: p.relative_to(root).as_posix().casefold())


def bundle_root(args: argparse.Namespace) -> Path:
    value = getattr(args, "root", None)
    root = Path(value).expanduser().resolve() if value else ROOT_DIR
    if not root.is_dir():
        raise ValueError(f"Code folder does not exist: {root}")
    return root


def bundle_target(root: Path, relative: str) -> Path:
    target = root.joinpath(*PurePosixPath(relative).parts)
    if not target.resolve().is_relative_to(root.resolve()):
        raise ValueError(f"Bundle path escapes the code folder: {relative}")
    return target


def bundle_pack(args: argparse.Namespace) -> int:
    root = bundle_root(args)
    paths = discover_sas(args.recursive, root)
    if not paths:
        raise ValueError("No SAS files found")
    hashes: list[str] = []
    chunks = [marker(BUNDLE_START, {
        "format": "sas-codebase-bundle", "version": FORMAT_VERSION,
        "generated_utc": utc_now(), "root_name": root.name,
        "recursive": bool(args.recursive), "file_count": len(paths),
    }, BUNDLE_END), "\n"]
    for path in paths:
        content = normalized(read_text(path))
        digest = sha256_text(content)
        hashes.append(digest)
        rel = path.relative_to(root).as_posix()
        chunks += [marker(FILE_START, {"path": rel, "sha256": digest, "chars": len(content)}, FILE_END),
                   content, f"{FILE_END}\n\n"]
    chunks.append(marker(FOOTER_START, {
        "file_count": len(paths), "manifest_sha256": sha256_text("".join(hashes)),
    }, FOOTER_END))
    output = root / (args.output or CODEBASE_FILE)
    output.write_text("".join(chunks), encoding="utf-8", newline="\n")
    print(f"Packed {len(paths)} SAS file(s): {output}")
    return 0


def parse_marker(line: str, prefix: str, suffix: str) -> dict[str, Any]:
    if not line.startswith(prefix) or not line.rstrip("\r\n").endswith(suffix):
        raise ValueError(f"Invalid bundle marker: expected {prefix}")
    raw = line[len(prefix): line.rfind(suffix)].strip()
    return json.loads(raw)


def parse_bundle(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    lines = path.read_text(encoding="utf-8-sig").splitlines(keepends=True)
    if not lines:
        raise ValueError("Bundle is empty")
    header = parse_marker(lines[0], BUNDLE_START, BUNDLE_END)
    files: list[dict[str, Any]] = []
    footer: dict[str, Any] | None = None
    i = 1
    while i < len(lines):
        line = lines[i]
        if line.startswith(FILE_START):
            meta = parse_marker(line, FILE_START, FILE_END)
            i += 1
            body: list[str] = []
            while i < len(lines) and lines[i].rstrip("\r\n") != FILE_END:
                body.append(lines[i]); i += 1
            if i >= len(lines):
                raise ValueError(f"Unclosed file block: {meta.get('path')}")
            files.append({"meta": meta, "content": normalized("".join(body))})
        elif line.startswith(FOOTER_START):
            footer = parse_marker(line, FOOTER_START, FOOTER_END)
        i += 1
    if footer is None:
        raise ValueError("Bundle footer is missing")
    return header, files, footer


def validate_bundle(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    header, files, footer = parse_bundle(path)
    errors: list[str] = []
    if header.get("format") != "sas-codebase-bundle" or header.get("version") != FORMAT_VERSION:
        errors.append("Unsupported header format/version")
    names: set[str] = set(); digests: list[str] = []
    for item in files:
        meta, content = item["meta"], item["content"]
        rel = str(meta.get("path", "")); pp = PurePosixPath(rel)
        if not rel or pp.is_absolute() or ".." in pp.parts:
            errors.append(f"Unsafe path: {rel!r}")
        key = rel.casefold()
        if key in names: errors.append(f"Duplicate path: {rel}")
        names.add(key)
        digest = sha256_text(content); digests.append(digest)
        if digest != meta.get("sha256"):
            item["edited"] = True
        else:
            item["edited"] = False
    if int(footer.get("file_count", -1)) != len(files):
        errors.append("Footer file count mismatch")
    original_manifest = sha256_text("".join(str(i["meta"].get("sha256", "")) for i in files))
    if footer.get("manifest_sha256") != original_manifest:
        errors.append("Footer manifest checksum mismatch")
    return files, errors


def bundle_verify(args: argparse.Namespace) -> int:
    root = bundle_root(args)
    path = root / (args.bundle or CODEBASE_FILE)
    files, errors = validate_bundle(path)
    edited = sum(bool(i["edited"]) for i in files)
    identical = different = missing = 0
    for item in files:
        target = bundle_target(root, item["meta"]["path"])
        if not target.exists(): missing += 1
        elif normalized(read_text(target)) == item["content"]: identical += 1
        else: different += 1
    print(f"Bundle: {path.name}")
    print(f"Files: {len(files)} | edited in bundle: {edited}")
    print(f"Compared with local source: {identical} identical, {different} different, {missing} missing")
    if errors:
        for error in errors: print(f"ERROR: {error}")
        return 1
    print("Bundle structure and paths are valid; safe to unpack.")
    return 0


def bundle_report(result: dict) -> None:
    """Optional observation hook; the standalone command owns all bundle writes."""
    print(f"Added: {result['added']}; changed: {result['changed']}; deleted: {result['deleted']} (files absent from bundle are retained).")


def bundle_unpack(args: argparse.Namespace) -> int:
    root = bundle_root(args)
    path = root / (args.bundle or CODEBASE_FILE)
    files, errors = validate_bundle(path)
    if errors: raise ValueError("; ".join(errors))
    changed: list[tuple[Path, str]] = []
    for item in files:
        target = bundle_target(root, item["meta"]["path"])
        if not target.exists() or normalized(read_text(target)) != item["content"]:
            changed.append((target, item["content"]))
    backup: Path | None = None
    existing = [p for p, _ in changed if p.exists()]
    if existing and not args.no_backup:
        backup = root / BACKUP_FOLDER / now_stamp()
        for source in existing:
            destination = backup / source.relative_to(root)
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, destination)
    for target, content in changed:
        target.parent.mkdir(parents=True, exist_ok=True)
        temp = target.with_name(target.name + f".{uuid.uuid4().hex}.tmp")
        temp.write_text(content, encoding="utf-8", newline="\n")
        os.replace(temp, target)
    print(f"Restored {len(changed)} changed/missing SAS file(s).")
    if backup: print(f"Backup: {backup}")
    bundle_report(dict(added=len(changed)-len(existing), changed=len(existing), deleted=0, backup=str(backup) if backup else ""))
    return 0


# ---------------------------------------------------------------------------
# EGP engine (ZIP/XML inspection and conservative extraction/repacking)
# ---------------------------------------------------------------------------


def choose_egp(value: str | None) -> Path:
    path = Path(value).expanduser() if value else find_one("*.egp", "EGP project")
    if not path.is_absolute(): path = (ROOT_DIR / path).resolve()
    if not path.is_file(): raise FileNotFoundError(path)
    return path


def egp_entries(path: Path) -> list[zipfile.ZipInfo]:
    with zipfile.ZipFile(path) as zf:
        return zf.infolist()


def embedded_sas_names(path: Path) -> list[str]:
    return sorted(i.filename for i in egp_entries(path) if i.filename.casefold().endswith(".sas"))


def egp_inspect(args: argparse.Namespace) -> int:
    path = choose_egp(args.project)
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist(); sas = [n for n in names if n.casefold().endswith(".sas")]
        print(f"Project: {path}")
        print(f"Archive entries: {len(names)} | embedded SAS members: {len(sas)}")
        for name in sas: print(f"  {name}")
        xmls = [n for n in names if n.casefold().endswith(".xml")]
        for name in xmls[:5]:
            try:
                root = ET.fromstring(zf.read(name))
                labels = []
                for node in root.iter():
                    tag = node.tag.rsplit("}", 1)[-1].casefold()
                    if tag in {"label", "name"} and node.text and node.text.strip():
                        labels.append(node.text.strip())
                if labels: print(f"  {name}: " + ", ".join(dict.fromkeys(labels[:12])))
            except Exception: pass
    return 0


# Extraction follows project labels and process-flow membership, as in MagicRoot.
# Keep this implementation self-contained: the CLI still ships as one Python file.
EGP_PROGRAMS = "01 SAS project files"
EGP_AUXILIARY = "02 Other files"
EGP_FLOWS = "03. By process flow"


def egp_archive_path(value: str) -> PurePosixPath:
    value = value.replace("\\", "/")
    path = PurePosixPath(value)
    if (not value or value.startswith("/") or PureWindowsPath(value).drive or
            ".." in path.parts or not path.parts or any(ord(c) < 32 for c in value)):
        raise ValueError("Unsafe path in EGP archive or manifest: " + repr(value))
    return path


def egp_component(value: str, fallback="Unnamed", limit=80) -> str:
    # Preserve readable labels, spaces and Unicode, while producing Windows filenames.
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip().rstrip(". ")
    if value in {"", ".", ".."}: value = fallback
    if re.fullmatch(r"(?i)(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?", value): value = "_" + value
    return value[:limit].rstrip(". ") or fallback


def egp_unique(value: str, used: set[str], suffix="") -> str:
    original = value
    number = 0
    while (value + suffix).casefold() in used:
        number += 1
        value = f"{original} ({number})"
    result = value + suffix
    used.add(result.casefold())
    return result


def egp_project_index(zf: zipfile.ZipFile, members: dict[str, zipfile.ZipInfo]):
    """Resolve classic PFD membership and newer container/parent relationships."""
    warnings = []
    xml = next((info for key, info in members.items() if key.casefold() == "project.xml"), None)
    if xml is None:
        return {}, {}, ["project.xml is missing; archive names and Unknown flows were used."]
    if xml.file_size > 16 * 1024**2: raise ValueError("EGP project.xml exceeds the 16 MB metadata limit.")
    raw = zf.read(xml)
    # Reject external/entity definitions; ordinary EG XML has no DTD.
    normalized_xml = raw.replace(b"\x00", b"").upper()
    if b"<!DOCTYPE" in normalized_xml or b"<!ENTITY" in normalized_xml:
        raise ValueError("EGP project.xml contains unsupported entity definitions.")
    try:
        root = ET.fromstring(raw)  # XML declarations/BOMs select UTF-8/UTF-16 correctly.
    except ET.ParseError:
        return {}, {}, ["project.xml could not be parsed; every embedded SAS file is still retained."]
    local = lambda node: node.tag.rsplit("}", 1)[-1].casefold()
    nodes = [node for node in root.iter() if local(node) == "element" and node.get("Type")]
    def text(node, names):
        # Prefer the element's own metadata over IDs inside PFD child processes.
        own = next((child for child in node if local(child) == "element" and not child.get("Type")), node)
        for name in names:
            for child in own.iter():
                if local(child) == name.casefold():
                    value = (child.text or child.get("ID") or "").strip()
                    if value: return value
        return ""
    flows, parents, programs, membership = {}, {}, {}, {}
    for node in nodes:
        kind = node.get("Type", "").rsplit(".", 1)[-1].casefold()
        identifier = text(node, ["ID"])
        label = text(node, ["Label", "Name"])
        if not identifier: continue
        parents[identifier] = text(node, ["Container", "ContainerID", "ParentID", "Parent", "OwnerID"])
        if kind in {"pfd", "processflow", "container"}:
            flows[identifier] = label or "Unnamed flow"
            if kind == "pfd":
                for process in node.iter():
                    if local(process) == "process":
                        reference = next(((child.text or "").strip() for child in process.iter() if local(child) == "id" and child.text), "")
                        if reference: membership.setdefault(reference, identifier)
        elif kind == "codetask":
            if identifier in programs: warnings.append("Duplicate program ID in project.xml: " + identifier)
            else: programs[identifier] = dict(label=label or "Unnamed_Code", parent=parents[identifier])
    for identifier, program in programs.items():
        parent = program["parent"]
        visited = set()
        while parent and parent not in visited:
            visited.add(parent)
            if parent in flows:
                program["flow"] = parent
                break
            parent = parents.get(parent, "")
        if not program.get("flow") and identifier in membership:
            program["flow"] = membership[identifier]
    return programs, flows, warnings


def egp_report(result: dict) -> None:
    """Optional observation hook for the local workbench."""
    print(f"Extracted {result['programs']} SAS program(s) in {result['flows']} process-flow folder(s): {result['path']}")
    for warning in result["warnings"]: print("WARNING: " + warning)


def egp_extract(args: argparse.Namespace) -> int:
    path = choose_egp(args.project)
    output = Path(args.output) if args.output else ROOT_DIR / path.stem
    if not output.is_absolute(): output = ROOT_DIR / output
    output = output.resolve()
    if output.exists(): raise FileExistsError(f"Output already exists; choose a different --output folder: {output}")
    with zipfile.ZipFile(path) as zf:
        infos = zf.infolist()
        if len(infos) > 20000 or sum(info.file_size for info in infos) > 2 * 1024**3:
            raise ValueError("EGP archive exceeds the 20,000-entry / 2 GB extraction limit.")
        members, raw_paths, used_paths = {}, {}, set()
        for info in infos:
            relative = egp_archive_path(info.orig_filename)
            key = str(relative).casefold()
            if key in used_paths: raise ValueError("Duplicate/colliding EGP archive path: " + info.filename)
            used_paths.add(key)
            if info.is_dir() or info.filename.endswith("\\"): continue
            if info.file_size > 256 * 1024**2: raise ValueError("EGP member exceeds the 256 MB limit: " + info.filename)
            if info.flag_bits & 1 or stat.S_ISLNK(info.external_attr >> 16):
                raise ValueError("Encrypted entries and symbolic links are not supported: " + info.filename)
            # Raw auxiliary paths are mapped safely too; the original ZIP remains intact.
            raw_paths[str(relative)] = PurePosixPath(*(egp_component(part, limit=100) for part in relative.parts))
            members[str(relative)] = info
        raw_keys = [str(relative).casefold() for relative in raw_paths.values()]
        if len(raw_keys) != len(set(raw_keys)): raise ValueError("EGP auxiliary filenames collide after Windows name sanitization.")
        sas = [key for key in members if key.casefold().endswith(".sas")]
        if not sas: raise ValueError("This EGP contains no embedded SAS source. Linked programs must be supplied separately.")
        programs, flows, warnings = egp_project_index(zf, members)
        lookup = {key.casefold(): key for key in sas}
        assignments, assigned = [], set()
        for identifier, program in programs.items():
            key = lookup.get((identifier.replace("\\", "/").rstrip("/") + "/code.sas").casefold())
            if key is None:
                warnings.append(f"No embedded code.sas for {program['label']} ({identifier}); it may be linked externally.")
                continue
            if key in assigned:
                warnings.append("More than one program references " + key + "; its source was retained once.")
                continue
            assigned.add(key)
            assignments.append((key, program["label"], program.get("flow")))
        for key in sas:
            if key not in assigned:
                leaf = PurePosixPath(key).stem
                if leaf.casefold() == "code" and len(PurePosixPath(key).parts) > 1: leaf = PurePosixPath(key).parts[-2]
                assignments.append((key, leaf, None))
                warnings.append("No program metadata matched " + key + "; its source was retained in an Unknown flow.")
        # Only a completely successful extraction becomes visible at the requested path.
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".pysas-egp-", dir=output.parent) as temp:
            staging = Path(temp) / "project"
            auxiliary = staging / EGP_AUXILIARY
            auxiliary.mkdir(parents=True)
            archive_copy = auxiliary / "original-egp.zip"
            if any(str(v).casefold() == "original-egp.zip" or str(v).casefold().startswith("original-egp.zip/") for v in raw_paths.values()):
                archive_copy = staging / "original-egp.zip"
            shutil.copy2(path, archive_copy)
            for key, info in members.items():
                target = auxiliary / raw_paths[key]
                if os.name == "nt" and max(len(str(output / EGP_AUXILIARY / raw_paths[key])), len(str(target))) > 240:
                    warnings.append("Long auxiliary path is preserved in original-egp.zip only: " + key)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info) as source, target.open("xb") as destination: shutil.copyfileobj(source, destination)
            manifest, aliases, hashes, records = {}, {}, {}, []
            flat_used, folder_used, grouped_used, flow_folders = set(), set(), {}, {}
            unknown = 0
            for key, label, flow in assignments:
                if flow:
                    flow_key = flow
                    flow_label = flows[flow]
                else:
                    flow_key = "unknown:" + key
                    flow_label = f"Unknown_{unknown}"
                    unknown += 1
                if flow_key not in flow_folders:
                    flow_limit = min(65, max(8, 240 - len(str(output / EGP_FLOWS)) - 50)) if os.name == "nt" else 65
                    flow_folders[flow_key] = egp_unique(egp_component(flow_label, "Unknown", flow_limit), folder_used)
                folder = flow_folders[flow_key]
                name = label[:-4] if label.casefold().endswith(".sas") else label
                name_limit = 90
                if os.name == "nt":
                    name_limit = min(name_limit, 240 - max(len(str(output / EGP_PROGRAMS)), len(str(output / EGP_FLOWS / folder))) - 12)
                    if name_limit < 8: raise ValueError("The extraction destination is too long. Choose a shorter --output folder.")
                name = egp_component(name, "Unnamed_Code", name_limit)
                flat_name = egp_unique(name, flat_used, ".sas")
                grouped_name = egp_unique(name, grouped_used.setdefault(folder, set()), ".sas")
                flat = PurePosixPath(EGP_PROGRAMS) / flat_name
                grouped = PurePosixPath(EGP_FLOWS) / folder / grouped_name
                # Preserve source bytes: extraction must never guess/re-encode SAS source.
                source_bytes = zf.read(members[key])
                for relative in (flat, grouped):
                    destination = staging / relative
                    destination.parent.mkdir(parents=True, exist_ok=True)
                    destination.write_bytes(source_bytes)
                canonical = str(flat)
                manifest[canonical] = members[key].filename
                aliases[canonical] = [str(grouped)]
                hashes[canonical] = hashlib.sha256(source_bytes).hexdigest()
                records.append(dict(path=canonical, by_flow=str(grouped), member=members[key].filename, label=label, flow=flows.get(flow) if flow else flow_label))
                print(f"  {label} -> {grouped}")
            (staging / ".pysas_egp_manifest.json").write_text(json.dumps(dict(
                format=2, source=str(path), members=manifest, aliases=aliases,
                hashes=hashes, programs=records, warnings=warnings,
            ), ensure_ascii=False, indent=2), encoding="utf-8")
            if output.exists(): raise FileExistsError("Output appeared during extraction; no files were replaced: " + str(output))
            staging.rename(output)
    egp_report(dict(path=str(output), programs=len(manifest), flows=len(flow_folders), warnings=warnings))
    return 0


def egp_pack(args: argparse.Namespace) -> int:
    source = Path(args.source)
    if not source.is_absolute(): source = ROOT_DIR / source
    source = source.resolve()
    template = choose_egp(args.template)
    manifest_path = source / ".pysas_egp_manifest.json"
    if not manifest_path.is_file(): raise FileNotFoundError("Extract this project with PySAS before repacking it.")
    saved = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest, aliases, hashes = saved.get("members"), saved.get("aliases", {}), saved.get("hashes", {})
    if not isinstance(manifest, dict) or not isinstance(aliases, dict) or not isinstance(hashes, dict): raise ValueError("Invalid EGP extraction manifest.")
    output = Path(args.output) if args.output else ROOT_DIR / f"{source.name}.egp"
    if not output.is_absolute(): output = ROOT_DIR / output
    output = output.resolve()
    if output.exists(): raise FileExistsError(f"Output already exists: {output}")
    replacements = {}
    with zipfile.ZipFile(template) as zin:
        available = set(zin.namelist())
        for leaf, member in manifest.items():
            if not isinstance(leaf, str) or not isinstance(member, str) or not member.casefold().endswith(".sas") or member not in available:
                raise ValueError("The extraction manifest does not match the selected EGP template.")
            egp_archive_path(member)
            extra = aliases.get(leaf, [])
            if not isinstance(extra, list) or any(not isinstance(value, str) for value in extra): raise ValueError("Invalid EGP editable aliases.")
            versions = []
            for relative in [leaf, *extra]:
                candidate = (source / egp_archive_path(relative)).resolve()
                if not candidate.is_relative_to(source): raise ValueError("An extracted program points outside its source folder.")
                if candidate.is_file(): versions.append((relative, candidate.read_bytes()))
            if not versions: continue  # Preserve the original member, matching legacy manifest behavior.
            baseline = hashes.get(leaf)
            if baseline is not None and (not isinstance(baseline, str) or not re.fullmatch(r"[0-9a-f]{64}", baseline)):
                raise ValueError("Invalid source hash in EGP manifest.")
            changed = [(relative, raw) for relative, raw in versions if baseline is None or hashlib.sha256(raw).hexdigest() != baseline]
            if len({raw for _, raw in changed}) > 1:
                raise ValueError("Conflicting edits in flat and by-flow copies of " + leaf + ". Keep one edit or make both copies identical.")
            if member in replacements: raise ValueError("Duplicate source mapping in EGP manifest: " + member)
            replacements[member] = changed[0][1] if changed else versions[0][1]
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_name(output.name + "." + uuid.uuid4().hex + ".tmp")
        try:
            with zipfile.ZipFile(temporary, "w") as zout:
                for info in zin.infolist(): zout.writestr(info, replacements.get(info.filename, zin.read(info)))
            if output.exists(): raise FileExistsError("Output appeared during repacking; no files were replaced.")
            temporary.rename(output)
        finally:
            temporary.unlink(missing_ok=True)
    print(f"Repacked {len(replacements)} SAS member(s): {output}")
    return 0


# ---------------------------------------------------------------------------
# Enterprise Guide automation and runner
# ---------------------------------------------------------------------------


VBS = r'''Option Explicit
Dim mode, projectPath, sasPath, programName, rowStart, rowEnd, logPath, codePath
Dim resultDir, tableManifest, tempPrefix, app, project, code, fso, stream, text, exitCode, sourceCode, selection
exitCode = 0
mode = WScript.Arguments(0)
projectPath = WScript.Arguments(1)
sasPath = WScript.Arguments(2)
programName = WScript.Arguments(3)
rowStart = CLng(WScript.Arguments(4))
rowEnd = CLng(WScript.Arguments(5))
logPath = WScript.Arguments(6)
codePath = WScript.Arguments(7)
resultDir = WScript.Arguments(8)
tableManifest = WScript.Arguments(9)
tempPrefix = WScript.Arguments(10)
Set fso = CreateObject("Scripting.FileSystemObject")

Function ReadAll(path)
  Dim s
  Set s = CreateObject("ADODB.Stream")
  s.Type = 2
  s.Charset = "utf-8"
  s.Open
  s.LoadFromFile path
  ReadAll = s.ReadText
  s.Close
End Function

Function RangeText(firstLine, lastLine)
  Dim first, last
  first = ""
  last = ""
  If firstLine > 0 Then first = CStr(firstLine)
  If lastLine > 0 Then last = CStr(lastLine)
  RangeText = ""
  If first <> "" Or last <> "" Then RangeText = first & ":" & last
End Function

Function SelectionLabel(ByVal firstLine, ByVal lastLine, ByVal section)
  SelectionLabel = "whole program"
  If section <> "" Then
    SelectionLabel = "section " & section
  ElseIf firstLine > 0 Or lastLine > 0 Then
    If firstLine = 0 Then firstLine = 1
    If lastLine = 0 Then lastLine = "end"
    SelectionLabel = "rows " & firstLine & " to " & lastLine
  End If
End Function

Sub CleanExit(message, status)
  WScript.Echo message
  On Error Resume Next
  project.Close
  app.Quit
  On Error GoTo 0
  WScript.Quit status
End Sub

Function FindProgram(requestedName)
    Dim item, matchCount, matchedItem
    matchCount = 0
    For Each item In project.CodeCollection
        If LCase(Trim(item.Name)) = LCase(Trim(requestedName)) Then
            matchCount = matchCount + 1
            Set matchedItem = item
        End If
    Next
    If matchCount = 0 Then CleanExit "ERROR: Program not found: " & requestedName, 14
    If matchCount > 1 Then CleanExit "ERROR: Program name is duplicated: " & requestedName, 29
    Set FindProgram = matchedItem
End Function

Function GetSelectedText(requestedName, requestedRange, requestedSection)
    Dim item, textValue
    Set item = FindProgram(requestedName)
    textValue = item.Text
    If requestedSection <> "" And requestedRange <> "" Then CleanExit "ERROR: Specify either section or row range, not both.", 23
    If requestedSection <> "" Then
        textValue = ExtractSection(textValue, requestedSection)
        WScript.Echo "Always/target section selected: " & requestedSection
    ElseIf requestedRange <> "" Then
        textValue = ExtractRange(textValue, requestedRange)
        WScript.Echo "Always/target range selected: " & requestedRange
    End If
    GetSelectedText = textValue
End Function

Function ExtractSection(textValue, requestedSection)
    Dim normalized, lines, i, startIndex, endIndex, startCount, endCount
    Dim startMarker, endMarker, selected
    normalized = Replace(textValue, vbCrLf, vbLf)
    normalized = Replace(normalized, vbCr, vbLf)
    lines = Split(normalized, vbLf)
    startMarker = LCase("* (please do not delete) section_start: " & requestedSection & ";")
    endMarker = LCase("* (please do not delete) section_end: " & requestedSection & ";")
    startIndex = -1
    endIndex = -1
    startCount = 0
    endCount = 0
    For i = 0 To UBound(lines)
        If LCase(Trim(lines(i))) = startMarker Then
            startCount = startCount + 1
            startIndex = i
        End If
        If LCase(Trim(lines(i))) = endMarker Then
            endCount = endCount + 1
            endIndex = i
        End If
    Next
    If startCount <> 1 Then CleanExit "ERROR: Expected exactly one section_start marker for: " & requestedSection, 24
    If endCount <> 1 Then CleanExit "ERROR: Expected exactly one section_end marker for: " & requestedSection, 25
    If endIndex <= startIndex Then CleanExit "ERROR: Section end must appear after section start: " & requestedSection, 26
    selected = ""
    For i = startIndex + 1 To endIndex - 1
        selected = selected & lines(i)
        If i < endIndex - 1 Then selected = selected & vbCrLf
    Next
    ExtractSection = selected
End Function

Function ExtractRange(textValue, requestedRange)
    Dim parts, normalized, lines, startLine, endLine, totalLines, i, selected
    parts = Split(requestedRange, ":")
    If UBound(parts) <> 1 Then CleanExit "ERROR: Invalid range. Expected 20:45, 20:, or :45", 15
    normalized = Replace(textValue, vbCrLf, vbLf)
    normalized = Replace(normalized, vbCr, vbLf)
    lines = Split(normalized, vbLf)
    totalLines = UBound(lines) + 1
    If Trim(parts(0)) = "" Then
        startLine = 1
    ElseIf IsNumeric(Trim(parts(0))) Then
        startLine = CLng(Trim(parts(0)))
    Else
        CleanExit "ERROR: Start line must be numeric.", 16
    End If
    If Trim(parts(1)) = "" Then
        endLine = totalLines
    ElseIf IsNumeric(Trim(parts(1))) Then
        endLine = CLng(Trim(parts(1)))
    Else
        CleanExit "ERROR: End line must be numeric.", 16
    End If
    If startLine < 1 Or endLine < startLine Then CleanExit "ERROR: Invalid line range: " & requestedRange, 17
    If endLine > totalLines Then CleanExit "ERROR: End line exceeds program length. Program lines: " & totalLines, 18
    selected = ""
    For i = startLine To endLine
        selected = selected & lines(i - 1)
        If i < endLine Then selected = selected & vbCrLf
    Next
    ExtractRange = selected
End Function

Function CleanName(value)
  Dim chars, i, c, answer
  chars = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-"
  answer = ""
  For i = 1 To Len(value)
    c = Mid(value, i, 1)
    If InStr(chars, c) > 0 Then answer = answer & c Else answer = answer & "_"
  Next
  If Len(answer) = 0 Then answer = "item"
  CleanName = Left(answer, 60)
End Function

Sub SaveOutputs(theCode)
  Dim i, item, outPath, manifest, itemName
  On Error Resume Next
  Err.Clear
  theCode.Log.SaveAs logPath
  If Err.Number <> 0 Then
    WScript.Echo "ERROR: log save failed: " & Err.Description
    If exitCode = 0 Then exitCode = 21
  Else
    WScript.Echo "Log saved: " & logPath
  End If
  Err.Clear
  If UCase(mode) <> "RUNFILE" Then
    On Error GoTo 0
    Exit Sub
  End If
  For i = 0 To theCode.Results.Count - 1
    Set item = theCode.Results.Item(i)
    itemName = CleanName(item.Name)
    outPath = resultDir & "\" & itemName & "_" & Right("000" & CStr(i + 1), 3) & ".html"
    item.SaveAs outPath
    If Err.Number = 0 Then WScript.Echo "HTML result saved: " & outPath Else WScript.Echo "WARNING: result save failed: " & Err.Description
    Err.Clear
  Next
  If Len(tableManifest) > 0 Then
    Set manifest = fso.OpenTextFile(tableManifest, 2, True)
    For i = 0 To theCode.OutputDatasets.Count - 1
      Set item = theCode.OutputDatasets.Item(i)
      itemName = CleanName(item.Name)
      outPath = tempPrefix & "_" & Right("000" & CStr(i + 1), 3) & ".xlsx"
      item.SaveAs outPath
      If Err.Number = 0 Then manifest.WriteLine itemName & vbTab & outPath Else WScript.Echo "WARNING: table export failed: " & Err.Description
      Err.Clear
    Next
    manifest.Close
  End If
  On Error GoTo 0
End Sub

On Error Resume Next
WScript.Echo "PYSAS_STAGE|opening|Starting Enterprise Guide automation"
Set app = CreateObject("SASEGObjectModel.Application.8.1")
If Err.Number <> 0 Then
  Err.Clear
  Set app = CreateObject("SASEGObjectModel.Application.7.1")
End If
If Err.Number <> 0 Then
  WScript.Echo "ERROR: Could not start SAS Enterprise Guide automation: " & Err.Description
  WScript.Quit 20
End If
WScript.Echo "PYSAS_STAGE|opening|Opening Enterprise Guide project"
Err.Clear
Set project = app.Open(projectPath, "")
If Err.Number <> 0 Then CleanExit "ERROR: Could not open project: " & Err.Description, 13
On Error GoTo 0

If UCase(mode) = "RUNFILE" Then
  If project.CodeCollection.Count = 0 Then CleanExit "ERROR: Template EGP needs a code object with a SAS server connection.", 33
  For Each sourceCode In project.CodeCollection
    Exit For
  Next
  text = ReadAll(sasPath)
  selection = "whole file"
Else
  Dim setupDoc, setupNode, initText, piece, section, stopProgram
  If Not fso.FileExists(sasPath) Then CleanExit "ERROR: Shared setup manifest not found: " & sasPath, 27
  Set setupDoc = CreateObject("MSXML2.DOMDocument.6.0")
  setupDoc.async = False
  If Not setupDoc.Load(sasPath) Then CleanExit "ERROR: Cannot read shared setup definitions", 28
  section = setupDoc.documentElement.getAttribute("section")
  stopProgram = setupDoc.documentElement.getAttribute("stop_program")
  If IsNull(section) Then section = ""
  initText = ""
  For Each setupNode In setupDoc.selectNodes("/setup/program")
    WScript.Echo "PYSAS_STAGE|preparing|Appending shared setup: " & setupNode.getAttribute("name") & " (" & SelectionLabel(CLng(setupNode.getAttribute("first")), CLng(setupNode.getAttribute("last")), setupNode.getAttribute("section")) & ")"
    piece = GetSelectedText(setupNode.getAttribute("name"), RangeText(CLng(setupNode.getAttribute("first")), CLng(setupNode.getAttribute("last"))), setupNode.getAttribute("section"))
    If initText <> "" Then initText = initText & vbCrLf
    initText = initText & "/* ALWAYS_RUN_ITEM: " & setupNode.getAttribute("name") & " */" & vbCrLf & piece
  Next
  selection = SelectionLabel(rowStart, rowEnd, section)
  WScript.Echo "PYSAS_STAGE|preparing|Appending task program: " & programName & " (" & selection & ")"
  piece = GetSelectedText(programName, RangeText(rowStart, rowEnd), section)
  text = "options iomlogautoflush;" & vbCrLf
  If stopProgram = "1" Then text = "options iomlogautoflush errorabend errorcheck=strict;" & vbCrLf
  If WScript.Arguments.Count > 11 Then text = text & ReadAll(WScript.Arguments(11)) & vbCrLf
  If initText <> "" Then text = text & "/* ALWAYS_RUN_INITIALIZATION_START */" & vbCrLf & initText & vbCrLf & "/* ALWAYS_RUN_INITIALIZATION_END */" & vbCrLf
  If WScript.Arguments.Count > 11 Then text = text & ReadAll(WScript.Arguments(11)) & vbCrLf
  text = text & "/* SCHEDULED_TARGET_START */" & vbCrLf & piece & vbCrLf & "/* SCHEDULED_TARGET_END */"
  Set sourceCode = FindProgram(programName)
End If

Set code = project.CodeCollection.Add
code.Text = text
' Preserve the source program's server, as in the working Rich-terminal 0.3.2.
code.Server = sourceCode.Server
If UCase(mode) = "RUNFILE" Then
  On Error Resume Next
  code.UseApplicationOptions = True
  On Error GoTo 0
End If
Dim saved
Set saved = CreateObject("ADODB.Stream")
saved.Type = 2
saved.Charset = "utf-8"
saved.Open
saved.WriteText code.Text
saved.SaveToFile codePath, 2
saved.Close
WScript.Echo "PYSAS_STAGE|prepared|Submitted code saved: " & codePath
WScript.Echo "PYSAS_STAGE|executing|Running: " & programName & " (" & selection & ")"
On Error Resume Next
Err.Clear
code.Run
If Err.Number <> 0 Then
  WScript.Echo "ERROR: Enterprise Guide reported an execution failure: " & Err.Description
  exitCode = 20
End If
On Error GoTo 0
WScript.Echo "PYSAS_STAGE|exporting|SAS returned; saving logs and results"
SaveOutputs code
If UCase(mode) = "RUNFILE" Then WScript.Echo "EG Results detected: " & code.Results.Count
WScript.Echo "PYSAS_STAGE|closing|Closing project"
project.Close
WScript.Echo "PYSAS_STAGE|closing|Closing Enterprise Guide"
app.Quit
WScript.Echo "PYSAS_STAGE|complete|Automation completed"
WScript.Quit exitCode
'''


def cscript_path() -> Path:
    windir = Path(os.environ.get("WINDIR", r"C:\Windows"))
    wow = windir / "SysWOW64" / "cscript.exe"
    return wow if wow.exists() else windir / "System32" / "cscript.exe"


def notify(title: str, message: str, error: bool = False, flash: bool = False) -> None:
    if os.name != "nt": return
    icon = "Error" if error else "Information"
    escaped_title = title.replace("'", "''"); escaped_message = message.replace("'", "''")
    script = (
        "Add-Type -AssemblyName System.Windows.Forms; Add-Type -AssemblyName System.Drawing; "
        "$n=New-Object System.Windows.Forms.NotifyIcon; "
        f"$n.Icon=[System.Drawing.SystemIcons]::{icon}; $n.BalloonTipTitle='{escaped_title}'; "
        f"$n.BalloonTipText='{escaped_message}'; $n.Visible=$true; $n.ShowBalloonTip(5000); "
        "Start-Sleep -Seconds 6; $n.Dispose()"
    )
    try:
        subprocess.Popen(["powershell", "-NoProfile", "-WindowStyle", "Hidden", "-Command", script],
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if flash:
            subprocess.Popen(["powershell", "-NoProfile", "-Command",
                              "$w=New-Object -ComObject WScript.Shell; $w.AppActivate($PID) | Out-Null"],
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except OSError:
        pass


def shared_work_settings(value):
    if not isinstance(value, dict):
        raise ValueError("Shared WORK settings must be an object.")
    enabled = value.get("enabled") is True
    path = str(value.get("path") or "").strip()
    libref = str(value.get("libref") or "VDI").strip().upper()
    if not re.fullmatch(r"[A-Z_][A-Z0-9_]{0,7}", libref) or libref in {"WORK", "SASHELP", "SASUSER", "USER"}:
        raise ValueError("Choose a library name of up to 8 letters, digits or underscores, such as VDI.")
    if len(path) > 2048 or any(ord(c) < 32 for c in path):
        raise ValueError("Enter a SAS server directory on one line (up to 2048 characters).")
    if enabled and not path:
        raise ValueError("Enter the WORK path from your open SAS session.")
    return {"enabled": enabled, "path": path, "libref": libref}


def shared_work_code():
    path = os.environ.get("PYSAS_SHARED_WORK_PATH", "")
    if not path: return ""
    config = shared_work_settings({"enabled": True, "path": path, "libref": os.environ.get("PYSAS_SHARED_WORK_LIBREF", "VDI")})
    literal = "'" + config["path"].replace("'", "''") + "'"
    lib = config["libref"]
    return ("/* PYSAS_SHARED_WORK_START */\n"
            f"data _null_; if libname('{lib}', {literal}) ne 0 then do; put 'ERROR: PySAS shared WORK is unavailable. Keep its owning session open and check the server path.'; abort abend; end; run;\n"
            f"options user={lib};\n"
            "/* PYSAS_SHARED_WORK_END */\n")


def compose_sas(source: Path, result_dir: Path, explicit_lib: Path | None = None,
                init_dir: Path | None = None, extra_init_dir: Path | None = None,
                parameters: Path | None = None, parameter_source: Path | None = None) -> str:
    parts = ["options iomlogautoflush;\n", shared_work_code(),
             f"%let PYSAS_RESULT_DIR=\"{result_dir.as_posix()}\";\n",
             "/* PYSAS_LIB_START */\n"]
    selected_parameters = (parameter_source or parameters).resolve() if parameters is not None else None
    for path in init_files(explicit_lib, init_dir, extra_init_dir):
        if selected_parameters is not None and path.resolve() == selected_parameters:
            continue  # A selected _*.sas parameter file belongs after initialization.
        parts += [f"/* PYSAS_INIT_FILE_START: {path.name} */\n", normalized(read_text(path)),
                  f"/* PYSAS_INIT_FILE_END: {path.name} */\n"]
    parts += ["/* PYSAS_LIB_END */\n", shared_work_code()]
    if parameters is not None:
        parts += ["/* PYSAS_PARAMETERS_START */\n", normalized(read_text(parameters)), "/* PYSAS_PARAMETERS_END */\n"]
    parts += [f"/* PYSAS_JOB_START: {source.name} */\n",
              normalized(read_text(source)), f"/* PYSAS_JOB_END: {source.name} */\n"]
    return "".join(parts)


def combine_workbooks(manifest_path: Path, output: Path) -> int:
    Workbook, load_workbook, _, _, _ = require_openpyxl()
    if not manifest_path.exists(): return 0
    rows = []
    for line in manifest_path.read_text(encoding="utf-8", errors="replace").splitlines():
        if "\t" in line:
            name, path = line.split("\t", 1); rows.append((name, Path(path)))
    if not rows: return 0
    final = Workbook(); final.remove(final.active); used: set[str] = set(); copied = 0
    for raw_name, temp_path in rows:
        if not temp_path.exists(): continue
        source = load_workbook(temp_path, data_only=False, read_only=False)
        for source_ws in source.worksheets:
            base = safe_name(raw_name or source_ws.title, 31)[:31] or "table"
            title = base; suffix = 2
            while title.casefold() in used:
                tail = f"_{suffix}"; title = base[:31-len(tail)] + tail; suffix += 1
            used.add(title.casefold()); target = final.create_sheet(title)
            for row in source_ws.iter_rows():
                for cell in row:
                    dest = target.cell(cell.row, cell.column, cell.value)
                    dest.number_format = cell.number_format
            copied += 1
        source.close()
        try: temp_path.unlink()
        except OSError: pass
    if copied: final.save(output)
    return copied


def detect_sas_error(log_path: Path) -> bool:
    if not log_path.exists(): return False
    return bool(re.search(r"(?mi)^\s*ERROR(?:\s+\d+-\d+)?:", read_text(log_path)))


def report_progress(run_dir: Path, phase: str, message: str) -> None:
    """The CLI and UI observe the same execution milestones."""
    with PRINT_LOCK:
        print(f"[{run_dir.name}] {message}", flush=True)


def forward_bridge_progress(path: Path, offset: int, pending: bytes, run_dir: Path):
    try:
        with path.open("rb") as source:
            source.seek(offset)
            chunk = source.read()
            offset = source.tell()
    except OSError:
        return offset, pending  # Observation must not interrupt an executing SAS process.
    lines = (pending + chunk).split(b"\n")
    pending = lines.pop()
    for raw in lines:
        line = raw.decode(text_encoding(raw), errors="replace").strip()
        if line.startswith("PYSAS_STAGE|") and line.count("|") >= 2:
            _, phase, message = line.split("|", 2)
            report_progress(run_dir, phase, message)
    return offset, pending


def execute_eg(mode: str, project: Path, sas_path: Path, program: str, row_start: int,
               row_end: int, run_dir: Path, tables: bool, cancel_file: str | Path | None = None, bridge: str | None = None) -> tuple[int, str]:
    project, sas_path, run_dir = project.resolve(), sas_path.resolve(), run_dir.resolve()
    logs = run_dir / "logs"; code_dir = run_dir / "code"; results = run_dir / "results"
    for folder in (logs, code_dir, results): folder.mkdir(parents=True, exist_ok=True)
    def cancellation_requested():
        return (run_dir / "_cancel.request").exists() or bool(cancel_file and Path(cancel_file).exists())
    if cancellation_requested():
        message = "Stopped by user before Enterprise Guide was started.\n"
        (run_dir / "console.txt").write_text(message, encoding="utf-8")
        return 130, message
    stem = safe_name(Path(program or sas_path.name).stem)
    log_path = logs / f"{stem}.log"; code_path = code_dir / f"{stem}.sas"
    manifest = run_dir / "_table_manifest.tsv" if tables else Path("")
    temp_prefix = run_dir / "_pysas_table"
    fd, vbs_name = tempfile.mkstemp(prefix="pysas_", suffix=".vbs")
    os.close(fd); vbs_path = Path(vbs_name)
    vbs_path.write_text(bridge or VBS, encoding="utf-16")
    command = [str(cscript_path()), "//nologo", str(vbs_path), mode, str(project), str(sas_path),
               program, str(row_start), str(row_end), str(log_path), str(code_path), str(results),
               str(manifest) if tables else "", str(temp_prefix)]
    shared_text = shared_work_code()
    if shared_text and mode.upper() not in {"RUNFILE", "RUNGRAPH"}:
        shared_path = run_dir / "_shared_work.sas"
        shared_path.write_text(shared_text, encoding="utf-8")
        command.append(str(shared_path))
    console_path = run_dir / "console.txt"
    rc = 127
    try:
        # A file cannot keep the parent waiting for EOF when EG leaves a child alive.
        with console_path.open("wb", buffering=0) as output:
            # Inherit the worker/terminal console and stdin, just as 0.3.2 did.
            process = subprocess.Popen(command, cwd=run_dir, stdout=output, stderr=subprocess.STDOUT,
                                       creationflags=subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0)
            cancelled = False
            offset, pending_output = 0, b""
            while process.poll() is None:
                offset, pending_output = forward_bridge_progress(console_path, offset, pending_output, run_dir)
                if cancellation_requested():
                    # Target only the automation process owned by this file, never all SAS/EG processes.
                    try:
                        if os.name == "nt":
                            stopped = subprocess.run(["taskkill.exe", "/PID", str(process.pid), "/T", "/F"],
                                                     capture_output=True, timeout=15, creationflags=subprocess.CREATE_NO_WINDOW)
                            if stopped.returncode and process.poll() is None:
                                raise OSError("Windows refused the stop request")
                        else:
                            process.kill()
                    except (OSError, subprocess.TimeoutExpired) as exc:
                        output.write(f"Stop request not completed: {exc}. Retrying while the file remains active.\n".encode("utf-8"))
                        output.flush()
                        time.sleep(3)
                        continue
                    cancelled = True
                    break
                time.sleep(.2)
            rc = process.wait()
            forward_bridge_progress(console_path, offset, pending_output, run_dir)
            if cancelled:
                rc = 130
                output.write(b"Stopped by user. The local automation process was terminated; verify remote SAS session state if needed.\n")
    except OSError as exc:
        with console_path.open("a", encoding="utf-8") as output:
            output.write(f"ERROR: Enterprise Guide automation: {exc}\n")
    finally:
        try: vbs_path.unlink()
        except OSError: pass
    console = re.sub(r"(?m)^PYSAS_STAGE\|[^|]+\|", "", read_text(console_path))
    if tables:
        combine_workbooks(manifest, results / f"{stem}_tables.xlsx")
        try: manifest.unlink()
        except OSError: pass
    return rc, console


def saved_runs(kind):
    """UI data root is optional; standalone CLI folder defaults stay compatible."""
    value = os.environ.get("PYSAS_RUNS_DIR")
    if value: return Path(value).expanduser().resolve() / kind
    return ROOT_DIR / "runs" if kind == "schedules" else ROOT_DIR / "runner" / "runs"


def run_job(source: Path, project: Path, tables: bool, explicit_lib: Path | None,
            notify_user: bool, runs_dir: Path | None = None, display_name: str | None = None,
            init_dir: Path | None = None, extra_init_dir: Path | None = None,
            parameters: Path | None = None) -> dict[str, Any]:
    if parameters is not None:
        parameters = Path(parameters).resolve()
        if parameters == source.resolve() or explicit_lib is not None and parameters == explicit_lib.resolve():
            raise ValueError("Choose a parameters file separate from the program and initialization override.")
        parameter_code = read_text(parameters)  # Read once; the archived copy is what will execute.
    logical_stem = re.sub(r"(?i)\.tables$", "", source.stem)
    base = safe_name(logical_stem)
    root = runs_dir or saved_runs("files")
    root.mkdir(parents=True, exist_ok=True)
    run_dir = root / f"{now_stamp()}__{base}"
    suffix = 2
    while run_dir.exists(): run_dir = root / f"{now_stamp()}__{base}_{suffix}"; suffix += 1
    run_dir.mkdir(parents=True)
    source_dir = run_dir / "source"
    source_dir.mkdir()
    shutil.copy2(source, source_dir / source.name)
    parameter_copy = None
    if parameters is not None:
        parameter_folder = run_dir / "parameters"
        parameter_folder.mkdir()
        parameter_copy = parameter_folder / parameters.name
        parameter_copy.write_text(parameter_code, encoding="utf-8")
        print(f"Parameters: {parameters.name} (after shared initialization, before {source.name})", flush=True)
    submitted = run_dir / "_submitted.sas"
    submitted.write_text(compose_sas(source, run_dir / "results", explicit_lib, init_dir, extra_init_dir,
                                    parameters=parameter_copy, parameter_source=parameters), encoding="utf-8")
    started = time.time()
    rc, console = execute_eg("RUNFILE", project, submitted, display_name or source.name, 0, 0, run_dir, tables)
    log_files = list((run_dir / "logs").glob("*.log"))
    sas_error = any(detect_sas_error(p) for p in log_files)
    if rc == 130: status = "CANCELLED"
    elif rc != 0: status = "FAILED"
    elif sas_error: status = "SAS_ERROR"
    else: status = "SUCCESS"
    elapsed = time.time() - started
    (run_dir / "status.txt").write_text(
        f"status={status}\nsource={source.name}\nstarted={datetime.fromtimestamp(started).isoformat(timespec='seconds')}\n"
        f"finished={datetime.now().isoformat(timespec='seconds')}\nelapsed_seconds={elapsed:.1f}\n",
        encoding="utf-8")
    if status == "SUCCESS":
        try: submitted.unlink()
        except OSError: pass
    if notify_user:
        label = "completed" if status == "SUCCESS" else ("completed with SAS errors" if status == "SAS_ERROR" else "runner failed")
        notify("PySAS", f"{source.name}: {label}", error=status != "SUCCESS")
    return {"name": source.name, "status": status, "run_dir": run_dir, "elapsed": elapsed, "console": console}


def runner_run(args: argparse.Namespace) -> int:
    source = Path(args.program).expanduser()
    if not source.is_absolute(): source = (Path.cwd() / source).resolve()
    if not source.is_file(): raise FileNotFoundError(source)
    project = choose_egp(args.template)
    tables = bool(args.tables or re.search(r"(?i)\.tables\.sas$", source.name))
    lib = Path(args.lib).resolve() if args.lib else None
    init_dir = Path(args.init_dir).expanduser().resolve() if getattr(args, "init_dir", None) else None
    if init_dir is not None and not init_dir.is_dir(): raise FileNotFoundError(init_dir)
    parameters = Path(args.parameters).expanduser().resolve() if getattr(args, "parameters", None) else None
    result = run_job(source, project, tables, lib, not args.no_notify, init_dir=init_dir, parameters=parameters)
    label = {"SUCCESS": "DONE", "SAS_ERROR": "DONE WITH SAS ERRORS", "FAILED": "RUNNER FAILED"}[result["status"]]
    print(f"{label}: {source.name}")
    print(result["run_dir"])
    return 0 if result["status"] == "SUCCESS" else 1


def format_elapsed(seconds: float) -> str:
    seconds = max(0, int(seconds)); hours, rem = divmod(seconds, 3600); minutes, secs = divmod(rem, 60)
    return f"{hours}:{minutes:02d}:{secs:02d}" if hours else f"{minutes:02d}:{secs:02d}"


def dashboard(running: dict[str, float], inbox: Path, runs: Path, max_ready: int = 3) -> str:
    lines = [f"PySAS {VERSION} watcher", "─" * 48, "", "Running now"]
    if running:
        for name, began in sorted(running.items(), key=lambda x: x[1]):
            lines.append(f"  ● {name:<30} {format_elapsed(time.time() - began)}")
    else: lines.append("  — idle —")
    ready = sorted((p for p in runs.iterdir() if p.is_dir()), key=lambda p: p.stat().st_mtime, reverse=True) if runs.exists() else []
    lines += ["", f"Ready to review ({len(ready)})"]
    for path in ready[:max_ready]: lines.append(f"  ✓ {path.name}")
    if len(ready) > max_ready: lines.append(f"  … and {len(ready) - max_ready} more in runner\\runs")
    queued = sum(1 for p in sas_files(inbox) if not p.name.startswith("_")) if inbox.exists() else 0
    if queued: lines += ["", f"Queued in inbox: {queued}"]
    lines += ["", f"Drop job .sas files into {inbox}. _*.sas files are shared initialization. Press Ctrl+C to stop."]
    return "\n".join(lines)


def runner_watch(args: argparse.Namespace) -> int:
    root = ROOT_DIR / "runner"
    inbox = Path(args.inbox).expanduser().resolve() if getattr(args, "inbox", None) else root / "inbox"
    # Claim on the inbox volume so custom/network inboxes can be moved atomically.
    claimed = root / "claimed" if inbox.resolve() == (root / "inbox").resolve() else inbox / ".pysas-claimed"
    runs = saved_runs("files")
    init_dir = Path(args.init_dir).expanduser().resolve() if getattr(args, "init_dir", None) else None
    if init_dir is not None and not init_dir.is_dir(): raise FileNotFoundError(init_dir)
    for path in (inbox, claimed, runs): path.mkdir(parents=True, exist_ok=True)
    project = choose_egp(args.template); lib = Path(args.lib).resolve() if args.lib else None
    max_workers = max(1, args.workers); running: dict[str, float] = {}; futures: dict[Any, tuple[str, Path]] = {}
    seen_stable: dict[Path, tuple[int, float]] = {}
    executor = concurrent.futures.ThreadPoolExecutor(max_workers=max_workers)
    ansi = sys.stdout.isatty()
    try:
        while True:
            for path in sorted(sas_files(inbox)):
                if path.name.startswith("_"): continue
                try: state = (path.stat().st_size, path.stat().st_mtime)
                except FileNotFoundError: continue
                if seen_stable.get(path) != state:
                    seen_stable[path] = state; continue
                if len(futures) >= max_workers: break
                claim_dir = claimed / uuid.uuid4().hex[:8]; claim_dir.mkdir()
                claimed_file = claim_dir / path.name
                try: path.replace(claimed_file)
                except (FileNotFoundError, PermissionError): continue
                seen_stable.pop(path, None)
                tables = bool(args.tables or re.search(r"(?i)\.tables\.sas$", claimed_file.name))
                future = executor.submit(run_job, claimed_file, project, tables, lib, not args.no_notify, runs, path.name, init_dir, inbox)
                futures[future] = (path.name, claim_dir); running[path.name] = time.time()
            for future in list(futures):
                if not future.done(): continue
                name, claim_dir = futures.pop(future); running.pop(name, None)
                try:
                    result = future.result(); print(f"\n{result['status']}: {name} -> {result['run_dir'].name}")
                except Exception as exc:
                    print(f"\nRUNNER FAILED: {name}: {exc}")
                    if not args.no_notify: notify("PySAS", f"{name}: runner failed", error=True)
                shutil.rmtree(claim_dir, ignore_errors=True)
            if ansi: print("\x1b[2J\x1b[H", end="")
            print(dashboard(running, inbox, runs), flush=True)
            time.sleep(max(0.5, args.poll))
    except KeyboardInterrupt:
        print("\nWatcher stopped. Running jobs will be allowed to finish.")
        return 130
    finally:
        executor.shutdown(wait=True, cancel_futures=False)


# ---------------------------------------------------------------------------
# Scheduler
# ---------------------------------------------------------------------------


def clean_int(value: Any, field: str, row: int) -> int | None:
    if value is None or str(value).strip() == "": return None
    number = float(value)
    if not number.is_integer(): raise ValueError(f"Excel row {row}: {field} must be an integer")
    return int(number)


def find_scheduler(value: str | None) -> Path:
    if value:
        path = Path(value); path = path if path.is_absolute() else ROOT_DIR / path
        return path.resolve()
    choices = sorted(p for p in ROOT_DIR.glob("*.xlsx") if not p.name.startswith("~$") and
                     p.name.casefold() != "run_summary.xlsx" and not p.stem.casefold().endswith("__next"))
    if len(choices) != 1: raise RuntimeError("Expected exactly one scheduler workbook; use --workbook")
    return choices[0]


def load_schedule(path: Path) -> list[dict[str, Any]]:
    _, load_workbook, _, _, _ = require_openpyxl()
    # Own the input handle: openpyxl row iterators can outlive early validation.
    with path.open("rb") as source:
        wb = load_workbook(source, data_only=True, read_only=True)
        try:
            ws = wb["Schedule"] if "Schedule" in wb.sheetnames else wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            headers = [str(v or "").strip().casefold() for v in next(rows, [])]
            named_headers = [h for h in headers if h]
            if len(set(named_headers)) != len(named_headers): raise ValueError("Scheduler contains duplicate column headers")
            missing = REQUIRED_COLUMNS.difference(headers)
            if missing: raise ValueError("Missing scheduler columns: " + ", ".join(sorted(missing)))
            col = {name: headers.index(name) for name in REQUIRED_COLUMNS}; tasks = []; ids: set[str] = set()
            for excel_row, values in enumerate(rows, 2):
                if not any(v is not None and str(v).strip() for v in values): continue
                def value(name: str): return values[col[name]] if col[name] < len(values) else None
                task_id = str(value("task_id") or "").strip()
                if not task_id: raise ValueError(f"Excel row {excel_row}: task_id is required")
                if task_id.casefold() in ids: raise ValueError(f"Excel row {excel_row}: duplicate task_id: {task_id}")
                ids.add(task_id.casefold())
                dependencies = [x.strip() for x in str(value("depends_on") or "").split(",") if x.strip()]
                tasks.append({
                    "task_id": task_id, "program": str(value("program") or "").strip(), "depends_on": dependencies,
                    "skip": truthy(value("skip")), "row_start": clean_int(value("row_start"), "row_start", excel_row),
                    "row_end": clean_int(value("row_end"), "row_end", excel_row), "section": str(value("section") or "").strip(),
                    "stop_process_on_error": truthy(value("stop_process_on_error")),
                    "stop_program_on_error": truthy(value("stop_program_on_error")),
                    "max_parallel": clean_int(value("max_parallel"), "max_parallel", excel_row),
                    "always_run": truthy(value("always_run")), "excel_row": excel_row,
                })
            if not tasks: raise ValueError("Scheduler contains no tasks")
            known = {t["task_id"].casefold() for t in tasks}
            for task in tasks:
                if not task["program"]: raise ValueError(f"{task['task_id']}: program is required")
                for field in ("row_start", "row_end"):
                    if task[field] == 0: task[field] = None
                    if task[field] is not None and task[field] < 1:
                        raise ValueError(f"{task['task_id']}: {field} must be blank, 0, or >= 1")
                if task["section"] and (task["row_start"] or task["row_end"]):
                    raise ValueError(f"{task['task_id']}: specify either section or row range, not both")
                if task["max_parallel"] is not None and task["max_parallel"] < 1:
                    raise ValueError(f"{task['task_id']}: max_parallel must be blank or at least 1")
                dependencies = [d.casefold() for d in task["depends_on"]]
                if len(dependencies) != len(set(dependencies)):
                    raise ValueError(f"{task['task_id']}: depends_on contains duplicates")
                absent = [d for d in task["depends_on"] if d.casefold() not in known]
                if absent: raise ValueError(f"{task['task_id']}: unknown dependencies: {', '.join(absent)}")
                if task["row_start"] and task["row_end"] and task["row_start"] > task["row_end"]:
                    raise ValueError(f"{task['task_id']}: row_start must not exceed row_end")
            dependencies = {t["task_id"].casefold(): {d.casefold() for d in t["depends_on"]} for t in tasks}
            waiting = set(dependencies)
            while waiting:
                ready = {key for key in waiting if not dependencies[key] & waiting}
                if not ready: raise ValueError("Circular dependency detected: " + ", ".join(sorted(waiting)))
                waiting -= ready
            return tasks
        finally:
            wb.close()


def describe_selection(task: dict[str, Any]) -> str:
    if task.get("section"):
        return f"section {task['section']}"
    if task.get("row_start") or task.get("row_end"):
        return f"rows {task.get('row_start') or 1} to {task.get('row_end') or 'end'}"
    return "whole program"


def scheduler_task(task: dict[str, Any], project: Path, task_root: Path) -> dict[str, Any]:
    if task.get("_graph_source"): return graph_task(task, project, task_root)
    task_dir = task_root / safe_name(task["task_id"]); task_dir.mkdir(parents=True, exist_ok=True)
    report_progress(task_dir, "preparing", f"Preparing {task['program']} ({describe_selection(task)})")
    task_project = task_dir / project.name; shutil.copy2(project, task_project)
    stop_program = task.get("stop_program_on_error", False) or any(t.get("stop_program_on_error") for t in task.get("_always_run", []))
    setup = ET.Element("setup", section=task.get("section", ""), stop_program="1" if stop_program else "0")
    for definition in task.get("_always_run", []):
        ET.SubElement(setup, "program", name=definition["program"],
                      first=str(definition["row_start"] or 0), last=str(definition["row_end"] or 0),
                      section=definition.get("section", ""))
    setup_path = task_dir / "shared_setup.xml"
    ET.ElementTree(setup).write(setup_path, encoding="utf-8", xml_declaration=True)
    report_progress(task_dir, "preparing", f"Shared setup: {len(task.get('_always_run', []))} definition(s) will be prepended")
    began = time.time()
    rc, console = execute_eg("RUNPROJECT", task_project, setup_path, task["program"],
                             task["row_start"] or 0, task["row_end"] or 0, task_dir, False,
                             cancel_file=task.get("_cancel_file"))
    logs = list((task_dir / "logs").glob("*.log")); sas_error = any(detect_sas_error(p) for p in logs)
    status = "CANCELLED" if rc == 130 else ("SAS_ERROR" if sas_error else ("FAILED" if rc else "SUCCESS"))
    evidence = console + "\n" + "\n".join(read_text(p) for p in logs)
    if any(pattern in evidence.casefold() for pattern in CONNECTION_PATTERNS): status = "CONNECTION_LOST"
    return {**task, "status": status, "elapsed": time.time() - began, "task_dir": task_dir,
            "message": console.strip().splitlines()[-1] if console.strip() else ""}


def report_task_result(result: dict[str, Any]) -> None:
    """Publish completion of a task that did not need an EG session."""
    with PRINT_LOCK:
        print(f"{result['task_id']}: {result['status']} · {result['message']}", flush=True)


def append_schedule_log_file(output, path: Path, console: bool = False) -> None:
    """Copy full text with bounded memory, including legacy Windows SAS encodings."""
    with path.open("rb") as source:
        encoding = text_encoding(source.read(4096))
    if encoding == "utf-8-sig":
        # An ASCII prefix does not distinguish UTF-8 from later cp1252 accents.
        try:
            with path.open(encoding=encoding) as source:
                while source.read(128 * 1024):
                    pass
        except UnicodeDecodeError:
            encoding = "cp1252"
    last = ""
    with path.open(encoding=encoding, errors="replace") as source:
        line_start = True
        while line := source.readline(128 * 1024):
            if console and line_start and line.startswith("PYSAS_STAGE|") and line.count("|") >= 2:
                line = line.split("|", 2)[2]
            output.write(line)
            line_start = line.endswith("\n")
            last = line
    if last and not last.endswith("\n"):
        output.write("\n")


def write_schedule_log(run_dir: Path, results: list[dict[str, Any]]) -> Path:
    """Combine task evidence in workbook order after a completed or stopped run."""
    target = run_dir / "schedule.log"
    fd, temporary_name = tempfile.mkstemp(prefix=".schedule-log-", suffix=".tmp", dir=run_dir)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as output:
            output.write(f"PySAS {VERSION} — Full schedule log\nSchedule: {run_dir.name}\n")
            output.write("Tasks are grouped in " + ("graph order" if (run_dir/"graph.json").is_file() else "workbook order") + "; parallel task durations overlap.\n")
            output.write("Includes available SAS logs and automation console output.\n\n")
            state = schedule_state(run_dir)
            if len(state["attempts"]) > 1:
                output.write(f"WARNING: This schedule ran in {len(state['attempts'])} parts. Pauses between parts are excluded.\n")
                output.write(f"Total active schedule time: {state['elapsed']:.1f} seconds.\n")
                if any(a.get("estimated") for a in state["attempts"]): output.write("Earlier timing is estimated from legacy task durations.\n")
                output.write("Each task below is its retained successful result or latest attempt. Earlier failed evidence is in attempts/.\n\n")
            for result in results:
                output.write(f"{result['task_id']}: {result['status']} — {result['program']} ({result.get('elapsed', 0):.1f}s)\n")
            for index, result in enumerate(results, 1):
                output.write("\n" + "=" * 80 + "\n")
                output.write(f"TASK {index}/{len(results)}: {result['task_id']} — {result['program']}\n")
                output.write(f"Selection: {describe_selection(result)}\nStatus: {result['status']}\nElapsed: {result.get('elapsed', 0):.1f} seconds\n")
                if result.get("message"):
                    output.write(f"Message: {result['message']}\n")
                folder = Path(result["task_dir"]) if result.get("task_dir") else None
                logs = sorted((folder / "logs").glob("*.log")) if folder else []
                if not logs:
                    note = "Shared setup is included in each target's SAS log." if result["status"] == "ALWAYS_RUN_DEFINITION" else "No SAS log was produced for this task."
                    output.write(note + "\n")
                evidence = [(path, False) for path in logs]
                if folder and (folder / "console.txt").is_file():
                    evidence.append((folder / "console.txt", True))
                for path, console in evidence:
                    output.write(f"\n--- {'Automation console' if console else 'SAS log'}: {path.name} ---\n")
                    try:
                        append_schedule_log_file(output, path, console=console)
                    except (OSError, ValueError) as exc:
                        output.write(f"\n[Could not read {path.name}: {exc}]\n")
            output.write("\n" + "=" * 80 + "\nEnd of full schedule log.\n")
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target


def write_summary(run_dir: Path, results: list[dict[str, Any]]) -> None:
    fields = ["task_id", "program", "status", "elapsed", "depends_on", "message"]
    with (run_dir / "run_summary.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields); writer.writeheader()
        for r in results: writer.writerow({k:",".join(r[k]) if k == "depends_on" else r.get(k, "") for k in fields})
    text = [f"PySAS {VERSION} schedule summary", ""]
    text += [f"{r['task_id']}: {r['status']} ({r.get('elapsed', 0):.1f}s)" for r in results]
    (run_dir / "run_summary.txt").write_text("\n".join(text) + "\n", encoding="utf-8")
    write_schedule_log(run_dir, results)
    if (run_dir / "graph.json").is_file():
        try: import openpyxl
        except ImportError: return
    Workbook, _, Font, PatternFill, get_column_letter = require_openpyxl()
    wb = Workbook(); ws = wb.active; ws.title = "Summary"; ws.append(fields)
    for cell in ws[1]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for r in results: ws.append([",".join(r[k]) if k == "depends_on" else r.get(k, "") for k in fields])
    for index, name in enumerate(fields, 1): ws.column_dimensions[get_column_letter(index)].width = max(12, len(name) + 2)
    wb.save(run_dir / "run_summary.xlsx")


def schedule_state(run_dir: Path) -> dict:
    path = run_dir / "schedule_state.json"
    if path.is_file():
        state = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(state.get("attempts"), list): raise ValueError("Invalid schedule attempt history")
        return state
    return {"attempts": [], "elapsed": 0.0}


def save_schedule_state(run_dir: Path, state: dict) -> None:
    temporary = run_dir / ".schedule-state.tmp"
    temporary.write_text(json.dumps(state, indent=2), encoding="utf-8")
    os.replace(temporary, run_dir / "schedule_state.json")


def report_schedule_state(run_dir: Path, state: dict) -> None:
    """Observer hook; execution remains in the single-file engine."""
    pass


@contextlib.contextmanager
def schedule_lock(run_dir: Path):
    lock_dir = run_dir.parent / ".schedule-locks"
    lock_dir.mkdir(exist_ok=True)
    with (lock_dir / (run_dir.name + ".lock")).open("a+b") as handle:
        try:
            if os.name == "nt":
                import msvcrt
                if handle.seek(0, 2) == 0: handle.write(b"0"); handle.flush()
                handle.seek(0); msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            raise ValueError("This schedule is already running in another process.") from exc
        yield


def schedule_run(args: argparse.Namespace) -> int:
    workbook = find_scheduler(args.workbook); project = choose_egp(args.project); tasks = load_schedule(workbook)
    if getattr(args, "_resume_dir", None):
        run_dir = Path(args._resume_dir)
    else:
        base_runs = saved_runs("schedules"); base_runs.mkdir(parents=True, exist_ok=True); run_dir = base_runs / f"{now_stamp()}__schedule"
        suffix = 2
        while run_dir.exists(): run_dir = base_runs / f"{now_stamp()}__schedule_{suffix}"; suffix += 1
        run_dir.mkdir(); shutil.copy2(workbook, run_dir / workbook.name); shutil.copy2(project, run_dir / project.name)
    workers = args.workers if args.workers is not None else max((t["max_parallel"] for t in tasks if t["max_parallel"] is not None), default=10)
    if workers < 1: raise ValueError("Parallel runs must be at least 1")
    with schedule_lock(run_dir):
        return execute_schedule(args, run_dir, project, tasks)


def execute_schedule(args, run_dir, project, tasks):
    state = schedule_state(run_dir)
    prior = {}
    if getattr(args, "_resume_dir", None):
        with (run_dir / "run_summary.csv").open(encoding="utf-8-sig", newline="") as source:
            prior = {r["task_id"].casefold(): r for r in csv.DictReader(source)}
        if not state["attempts"]:
            known = getattr(args, "prior_elapsed", None)
            estimate = max((float(r.get("elapsed") or 0) for r in prior.values()), default=0)
            seconds = known if known is not None else estimate
            if not math.isfinite(seconds) or seconds < 0: raise ValueError("Prior elapsed time must be non-negative and finite")
            state = {"attempts": [{"number": 1, "elapsed": seconds, "status": "IMPORTED", "estimated": known is None}], "elapsed": seconds}
        archive = run_dir / "attempts" / f"part-{len(state['attempts']):03d}"
        archive.mkdir(parents=True, exist_ok=True)
        for name in ("run_summary.csv", "run_summary.xlsx", "run_summary.txt", "schedule.log", "status.txt"):
            if (run_dir / name).is_file(): shutil.copy2(run_dir / name, archive / name)
        # Preserve failed evidence separately; the current tasks folder has one result per task.
        for task in tasks:
            key = task["task_id"].casefold()
            folder = run_dir / "tasks" / safe_name(task["task_id"])
            if not task["always_run"] and prior.get(key, {}).get("status") not in {"SUCCESS", "SKIPPED_PREVIOUS"} and folder.exists():
                destination = archive / "tasks" / folder.name
                destination.parent.mkdir(parents=True, exist_ok=True)
                if destination.exists(): raise ValueError("Prior attempt archive already contains this task; inspect it before continuing")
                folder.rename(destination)
        (run_dir / "_cancel.request").unlink(missing_ok=True)
    prior_elapsed = float(state.get("elapsed", 0))
    began = time.time()
    state["attempts"].append({"number": len(state["attempts"])+1, "started": began, "status": "RUNNING", "elapsed": 0})
    save_schedule_state(run_dir, state)
    report_schedule_state(run_dir, state)
    cancel_file = Path(args.cancel_file).resolve() if getattr(args, "cancel_file", None) else run_dir / "_cancel.request"
    definitions = [t for t in tasks if t["always_run"] and not t["skip"]]
    pending = {t["task_id"].casefold(): {**t, "_always_run": definitions, "_cancel_file": str(cancel_file)} for t in tasks if not t["always_run"]}
    results = [{**t, "status": "SKIPPED_SUCCESS" if t["skip"] else "ALWAYS_RUN_DEFINITION",
                "elapsed": 0.0, "message": "Marked skip" if t["skip"] else "Prepended before each program in " + ("graph order" if (run_dir/"graph.json").is_file() else "workbook order")}
               for t in tasks if t["always_run"]]
    for key, task in list(pending.items()):
        previous = prior.get(key, {})
        if previous.get("status") in {"SUCCESS", "SKIPPED_PREVIOUS"}:
            kept = {**task, "status": previous["status"], "elapsed": float(previous.get("elapsed") or 0),
                    "message": "Completed in an earlier part; original output retained" if (run_dir / "tasks" / safe_name(task["task_id"])).is_dir() else "Previously completed; its original output folder is missing", "task_dir": run_dir / "tasks" / safe_name(task["task_id"])}
            results.append(kept); pending.pop(key); report_task_result(kept)
    satisfied = {t["task_id"].casefold() for t in tasks if t["always_run"]}
    satisfied.update(r["task_id"].casefold() for r in results if r["status"] in {"SUCCESS", "SKIPPED_PREVIOUS"})
    failed: set[str] = set(); stop = False; cancelled = False
    max_workers = args.workers if args.workers is not None else max((t["max_parallel"] for t in tasks if t["max_parallel"] is not None), default=10)
    if max_workers < 1: raise ValueError("Parallel runs must be at least 1")
    executor = concurrent.futures.ThreadPoolExecutor(max_workers=max_workers); active: dict[Any, dict[str, Any]] = {}
    submitted_at: dict[str, float] = {}
    def resolve_pending(key, task, status, message):
        result = {**task, "status": status, "elapsed": 0.0, "message": message}
        results.append(result)
        (satisfied if status in FINAL_OK else failed).add(key)
        pending.pop(key)
        report_task_result(result)
    try:
        while pending or active:
            cancelled = cancelled or cancel_file.exists()
            pending_before = len(pending)
            for key, task in list(pending.items()):
                # Check the run-wide request before every launch, including skipped rows.
                cancelled = cancelled or cancel_file.exists()
                if cancelled:
                    resolve_pending(key, task, "CANCELLED", "Schedule stopped before this task started")
                    continue
                deps = {d.casefold() for d in task["depends_on"]}
                if stop:
                    resolve_pending(key, task, "STOPPED_ON_ERROR", "Not started after stop-on-error")
                    continue
                if deps & failed:
                    resolve_pending(key, task, "BLOCKED_DEPENDENCY", "Dependency failed")
                    continue
                # A skipped program is a dependency barrier, not an already completed job.
                # Resolving each barrier after its parents preserves arbitrary skipped chains.
                if not deps.issubset(satisfied): continue
                if task["skip"]:
                    resolve_pending(key, task, "SKIPPED_SUCCESS", "Skipped after dependencies completed")
                    continue
                if len(active) >= max_workers: continue
                submitted_at[key] = time.time()
                print(f"Launching {task['task_id']}: {task['program']} ({describe_selection(task)})", flush=True)
                future = executor.submit(scheduler_task, task, project, run_dir / "tasks")
                active[future] = task; pending.pop(key)
            if not active:
                if pending and len(pending) < pending_before:
                    continue
                if pending:
                    for key, task in list(pending.items()):
                        results.append({**task, "status": "BLOCKED_DEPENDENCY", "elapsed": 0.0,
                                        "message": "Circular/unresolved dependency"}); failed.add(key); pending.pop(key)
                break
            done, _ = concurrent.futures.wait(active, timeout=0.5, return_when=concurrent.futures.FIRST_COMPLETED)
            for future in done:
                task = active.pop(future); key = task["task_id"].casefold()
                try:
                    result = future.result()
                except Exception as exc:
                    result = {**task, "status": "FAILED", "elapsed": time.time() - submitted_at[key],
                              "message": str(exc), "task_dir": run_dir / "tasks" / safe_name(task["task_id"])}
                results.append(result)
                if result["status"] == "SUCCESS": satisfied.add(key)
                else:
                    failed.add(key)
                    if result["status"] == "CONNECTION_LOST" or task["stop_process_on_error"] or any(t.get("stop_process_on_error") for t in definitions): stop = True
                print(f"{task['task_id']}: {result['status']}")
    finally: executor.shutdown(wait=True)
    order = {t["task_id"].casefold(): i for i, t in enumerate(tasks)}; results.sort(key=lambda r: order[r["task_id"].casefold()])
    cancelled = cancelled or cancel_file.exists()
    errors = [r for r in results if r["status"] not in FINAL_OK]
    finished = time.time()
    status = "STOPPED" if cancelled else ("FAILED" if errors else "SUCCESS")
    state["attempts"][-1].update(finished=finished, elapsed=max(0, finished-began), status=status)
    state["elapsed"] = prior_elapsed + state["attempts"][-1]["elapsed"]
    save_schedule_state(run_dir, state)
    (run_dir / "status.txt").write_text(
        f"status={status}\nstarted={datetime.fromtimestamp(state['attempts'][0].get('started', began)).isoformat()}\nelapsed_seconds={state['elapsed']:.3f}\n", encoding="utf-8")
    write_summary(run_dir, results)
    report_schedule_state(run_dir, state)
    if not args.no_notify: notify("PySAS schedule", "Stopped by user" if cancelled else ("Completed successfully" if not errors else f"Completed with {len(errors)} error(s)"),
                                  error=bool(errors) and not cancelled, flash=bool(errors) and not cancelled)
    print(f"Schedule {'stopped' if cancelled else 'complete'}: {run_dir}")
    return 130 if cancelled else (1 if errors else 0)


def schedule_continue(args: argparse.Namespace) -> int:
    previous = Path(args.run_dir); previous = previous if previous.is_absolute() else ROOT_DIR / previous
    summaries = previous / "run_summary.csv"
    if not summaries.is_file(): raise FileNotFoundError(summaries)
    workbook = next((p for p in previous.glob("*.xlsx") if p.name != "run_summary.xlsx"), None)
    project = next(iter(previous.glob("*.egp")), None)
    if not workbook or not project: raise FileNotFoundError("Continuation folder must contain its scheduler workbook and EGP")
    forwarded = argparse.Namespace(workbook=str(workbook), project=str(project), workers=args.workers, no_notify=args.no_notify,
                                   cancel_file=getattr(args, "cancel_file", None), _resume_dir=previous,
                                   prior_elapsed=getattr(args, "prior_elapsed", None))
    return schedule_run(forwarded)


# ---------------------------------------------------------------------------
# JSON dependency graphs (Excel schedules keep their existing execution path)
# ---------------------------------------------------------------------------
GRAPH_FORMAT = "pysas-graph-1"
GRAPH_LIMIT = 2 * 1024**2
GRAPH_SOURCE_LIMIT = 16 * 1024**2


def validate_graph(value, *, runnable=False):
    if not isinstance(value, dict) or value.get("format") != GRAPH_FORMAT:
        raise ValueError("Choose a PySAS graph JSON file (pysas-graph-1).")
    def text(v, label, limit=2048):
        if not isinstance(v, str) or len(v) > limit or any(ord(c) < 32 for c in v):
            raise ValueError(label + " is invalid.")
        return v.strip()
    graph_id = text(value.get("id", ""), "Graph ID", 64)
    if graph_id and (not re.fullmatch(r"[a-zA-Z0-9_-]+", graph_id) or egp_component(graph_id) != graph_id): raise ValueError("Invalid graph ID.")
    name = text(value.get("name", "Untitled graph"), "Graph name", 120) or "Untitled graph"
    workers = value.get("workers", 10)
    if type(workers) is not int or not 1 <= workers <= 32: raise ValueError("Choose 1 to 32 parallel programs.")
    connection = text(value.get("connection", ""), "Connection EGP")
    nodes, setup, seen = [], [], set()
    for field, output, maximum in (("setup", setup, 100), ("nodes", nodes, 200)):
        entries = value.get(field, [])
        if not isinstance(entries, list) or len(entries) > maximum: raise ValueError(f"Too many {field} entries (maximum {maximum}).")
        for node in entries:
            if not isinstance(node, dict): raise ValueError("Each program must be an object.")
            identifier = text(node.get("id"), "Program ID", 60)
            if not re.fullmatch(r"[a-zA-Z0-9][a-zA-Z0-9_-]{0,59}", identifier) or egp_component(identifier) != identifier:
                raise ValueError("Program IDs need letters, digits, underscores or hyphens, and a valid Windows filename.")
            if identifier.casefold() in seen: raise ValueError("Duplicate program ID: " + identifier)
            seen.add(identifier.casefold())
            source = node.get("source")
            if not isinstance(source, dict) or source.get("type") not in {"file", "egp"}: raise ValueError(identifier + ": choose a SAS file or embedded EGP program.")
            path = text(source.get("path", ""), "Source path")
            if not path: raise ValueError(identifier + ": source path is required.")
            cleaned = dict(type=source['type'], path=path)
            if source['type'] == 'egp':
                cleaned['program_id'] = text(source.get('program_id', ''), 'EGP program ID', 200)
                if not cleaned['program_id']: raise ValueError(identifier + ': choose an embedded EGP program.')
            item = dict(id=identifier, name=text(node.get('name', identifier), 'Program name', 120) or identifier, source=cleaned)
            for flag, default in (('skip', False), ('stop_process_on_error', True), ('stop_program_on_error', False), ('tables', False)):
                v = node.get(flag, default)
                if type(v) is not bool: raise ValueError(identifier + ': '+flag+' must be true or false.')
                item[flag] = v
            item['section'] = text(node.get('section', ''), 'Section name', 200)
            for bound in ('row_start', 'row_end'):
                v = node.get(bound)
                if v == 0: v = None
                if v is not None and (type(v) is not int or v < 1): raise ValueError(identifier + ': row bounds must be positive integers.')
                item[bound] = v
            if item['section'] and (item['row_start'] or item['row_end']): raise ValueError(identifier + ': choose a section or row range, not both.')
            if item['row_start'] and item['row_end'] and item['row_start'] > item['row_end']: raise ValueError(identifier + ': row_start exceeds row_end.')
            for coordinate in ('x', 'y'):
                v = node.get(coordinate, 0)
                if type(v) not in (int, float) or not math.isfinite(v) or not 0 <= v <= 20000: raise ValueError('Canvas positions must be between 0 and 20,000.')
                item[coordinate] = round(v, 2)
            output.append(item)
    if runnable and not any(not n['skip'] for n in nodes): raise ValueError('Add at least one enabled program below shared setup.')
    edges = value.get('edges', [])
    if not isinstance(edges, list) or len(edges) > 2000: raise ValueError('Graph has too many dependencies (maximum 2,000).')
    known = {n['id'] for n in nodes}; pairs = set(); links = []
    for edge in edges:
        if not isinstance(edge, dict): raise ValueError('Each dependency needs from and to IDs.')
        first, last = edge.get('from'), edge.get('to')
        if not isinstance(first, str) or not isinstance(last, str) or first not in known or last not in known: raise ValueError('Dependencies must link existing program cards, not shared setup.')
        if (first, last) in pairs: raise ValueError('Duplicate dependency.')
        if first == last: raise ValueError('A program cannot depend on itself.')
        pairs.add((first, last)); links.append(dict(**{'from':first}, to=last))
    waiting = set(known); dependencies = {k:{a for a,b in pairs if b==k} for k in known}
    while waiting:
        ready = {k for k in waiting if not dependencies[k] & waiting}
        if not ready: raise ValueError('Circular dependency detected. Remove a link before saving or running.')
        waiting -= ready
    return dict(format=GRAPH_FORMAT, id=graph_id, name=name, connection=connection, workers=workers, setup=setup, nodes=nodes, edges=links)


def load_graph(path, *, runnable=False):
    path = Path(path)
    if path.stat().st_size > GRAPH_LIMIT: raise ValueError('Graph JSON exceeds 2 MB.')
    return validate_graph(json.loads(path.read_text(encoding='utf-8-sig')), runnable=runnable)


def graph_source_path(value, base, suffix):
    path = Path(value).expanduser()
    path = (base/path).resolve() if not path.is_absolute() else path.resolve()
    if not path.is_file() or path.suffix.casefold() != suffix: raise ValueError('Source is missing or has the wrong type: '+str(path))
    return path


def graph_egp_programs(path):
    """Inspect embedded source and labels without opening Enterprise Guide."""
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos)>20000 or sum(i.file_size for i in infos)>2*1024**3: raise ValueError('EGP exceeds the extraction limits.')
        members = {}; seen_members=set()
        for info in infos:
            key = str(egp_archive_path(info.orig_filename))
            if info.flag_bits & 1 or stat.S_ISLNK(info.external_attr >> 16): raise ValueError('Encrypted/symbolic EGP members are unsupported.')
            if info.file_size > 256*1024**2: raise ValueError('EGP member exceeds 256 MB.')
            if key.casefold() in seen_members: raise ValueError('Duplicate EGP archive path.')
            seen_members.add(key.casefold());members[key] = info
        programs, flows, warnings = egp_project_index(archive, members)
        lookup = {k.casefold():i for k,i in members.items()}
        counts = {}
        for program in programs.values(): counts[program['label'].casefold()] = counts.get(program['label'].casefold(), 0)+1
        result = []
        for identifier, program in programs.items():
            info = lookup.get((identifier+'/code.sas').casefold())
            result.append(dict(id=identifier, name=program['label'], flow=flows.get(program.get('flow'), 'Unassigned'),
                               member=info.filename if info else None, duplicate=counts[program['label'].casefold()]>1))
        return result, warnings


def select_graph_text(text, node):
    lines = text.replace('\r\n','\n').replace('\r','\n').split('\n')
    section = node.get('section')
    if section:
        start = f'* (please do not delete) section_start: {section};'.casefold()
        end = f'* (please do not delete) section_end: {section};'.casefold()
        starts = [i for i,line in enumerate(lines) if line.strip().casefold()==start]
        ends = [i for i,line in enumerate(lines) if line.strip().casefold()==end]
        if len(starts)!=1 or len(ends)!=1 or starts[0]>=ends[0]: raise ValueError(node['name']+': section markers are missing, duplicated or out of order.')
        return '\n'.join(lines[starts[0]+1:ends[0]])
    first, last = node.get('row_start') or 1, node.get('row_end') or len(lines)
    if last>len(lines) or first>last: raise ValueError(node['name']+': row selection exceeds the source length.')
    return '\n'.join(lines[first-1:last])


def report_graph_plan(tasks):
    """Observation hook; the dependency scheduler remains the execution owner."""
    pass


def graph_tasks(run_dir):
    raw = json.loads((run_dir/'graph.json').read_text(encoding='utf-8'))
    graph = validate_graph(raw, runnable=True)
    evidence = json.loads((run_dir/'graph-sources.json').read_text(encoding='utf-8'))
    def asset(relative):
        path = (run_dir/egp_archive_path(relative)).resolve()
        if not path.is_relative_to(run_dir.resolve()) or not path.is_file(): raise ValueError('Graph snapshot is missing or points outside its run.')
        if path in checked: return path
        expected = evidence['files'].get(str(egp_archive_path(relative)))
        digest=hashlib.sha256()
        with path.open('rb') as stream:
            for chunk in iter(lambda:stream.read(1024*1024),b''):digest.update(chunk)
        if not expected or digest.hexdigest()!=expected: raise ValueError('Graph snapshot changed; start a new graph run instead of continuing edited source.')
        checked[path]=True
        return path
    tasks=[]
    checked={}
    for field in ('setup','nodes'):
        for node in graph[field]:
            saved=evidence['nodes'].get(node['id'])
            if not saved: raise ValueError('Graph source manifest is incomplete.')
            source=asset(saved['selected']); project=asset(saved['project']) if saved.get('project') else None
            tasks.append(dict(task_id=node['id'], program=node['name'], section=node['section'], row_start=node['row_start'], row_end=node['row_end'],
                              skip=node['skip'], always_run=field=='setup', max_parallel=graph['workers'], depends_on=[e['from'] for e in graph['edges'] if e['to']==node['id']],
                              stop_process_on_error=node['stop_process_on_error'], stop_program_on_error=node['stop_program_on_error'],
                              _graph_source=dict(path=str(source), project=str(project) if project else '', program=saved.get('egp_name',''), tables=node['tables'])))
    report_graph_plan(tasks)
    return graph, tasks


def prepare_graph_run(path, workers=None):
    path=Path(path).resolve(); graph=load_graph(path, runnable=True); base=path.parent
    if workers is not None: graph=validate_graph({**graph,'workers':workers}, runnable=True)
    root=saved_runs('schedules');root.mkdir(parents=True,exist_ok=True)
    target=root/(now_stamp()+'__graph_'+safe_name(graph['name'],35));number=2
    while target.exists():target=root/(now_stamp()+'__graph_'+safe_name(graph['name'],35)+'_'+str(number));number+=1
    with tempfile.TemporaryDirectory(prefix='.graph-preparing-',dir=root) as temporary:
        stage=Path(temporary)/'run';sources=stage/'sources';sources.mkdir(parents=True)
        files={};records={};projects={};indexes={};connection_origin=graph['connection']
        def keep_bytes(name,raw):
            (stage/name).write_bytes(raw);files[name]=hashlib.sha256(raw).hexdigest();return name
        def project_copy(value):
            project=graph_source_path(value,base,'.egp')
            if project in projects:return projects[project]
            if project.stat().st_size>512*1024**2:raise ValueError('Connection EGP exceeds 512 MB.')
            name='sources/project_'+str(len(projects))+'.egp'
            shutil.copy2(project,stage/name)
            digest=hashlib.sha256()
            with (stage/name).open('rb') as stream:
                for chunk in iter(lambda:stream.read(1024*1024),b''):digest.update(chunk)
            files[name]=digest.hexdigest();projects[project]=name;return name
        connection=project_copy(graph['connection']) if graph['connection'] else ''
        graph['connection']=connection
        for field in ('setup','nodes'):
            for node in graph[field]:
                print('Snapshotting '+node['name']+' ('+describe_selection(node)+')',flush=True)
                source=node['source']; origin=dict(source);egp_name='';project='';program=None
                if source['type']=='file':
                    file=graph_source_path(source['path'],base,'.sas')
                    if file.stat().st_size>GRAPH_SOURCE_LIMIT:raise ValueError(node['name']+': SAS source exceeds 16 MB.')
                    raw=file.read_bytes();source['path']=keep_bytes('sources/'+node['id']+'.sas',raw)
                    text=raw.decode(text_encoding(raw),errors='strict');project=connection
                else:
                    original=graph_source_path(source['path'],base,'.egp');project=project_copy(str(original))
                    if project not in indexes:indexes[project]=graph_egp_programs(stage/project)[0]
                    program=next((p for p in indexes[project] if p['id']==source['program_id']),None)
                    if not program or not program['member']:raise ValueError(node['name']+': no readable embedded source; linked EGP programs must be supplied as SAS files.')
                    if field=='nodes' and program['duplicate']:raise ValueError(node['name']+': duplicate EGP program names are ambiguous for server selection. Give the program a unique name in EG first.')
                    with zipfile.ZipFile(stage/project) as archive:
                        info=archive.getinfo(program['member'])
                        if info.file_size>GRAPH_SOURCE_LIMIT:raise ValueError(node['name']+': SAS source exceeds 16 MB.')
                        raw=archive.read(info)
                    text=raw.decode(text_encoding(raw),errors='strict');source['path']=project;egp_name=program['name']
                if field=='nodes' and not project and not node['skip']:raise ValueError('Choose a connection EGP for linked SAS programs.')
                selected=select_graph_text(text,node)
                selected_path=keep_bytes('sources/'+node['id']+'.selected.sas',selected.encode('utf-8'))
                records[node['id']]=dict(selected=selected_path, project=project, egp_name=egp_name,
                                         origin=origin,
                                         source_hash=hashlib.sha256(raw).hexdigest(),selected_hash=hashlib.sha256(selected.encode('utf-8')).hexdigest())
        (stage/'graph.json').write_text(json.dumps(graph,ensure_ascii=False,indent=2),encoding='utf-8')
        (stage/'graph-sources.json').write_text(json.dumps(dict(format=GRAPH_FORMAT,files=files,nodes=records,connection_origin=connection_origin,graph_hash=hashlib.sha256((stage/'graph.json').read_bytes()).hexdigest()),indent=2),encoding='utf-8')
        stage.rename(target)
    return target


def graph_bridge():
    # This branch is used only by graph EGP nodes. Legacy RUNPROJECT/RUNFILE
    # behavior and their server-selection rules stay exactly as before.
    marker='If UCase(mode) = "RUNFILE" Then\n  If project.CodeCollection.Count'
    replacement='If UCase(mode) = "RUNGRAPH" Then\n  Set sourceCode = FindProgram(programName)\n  text = ReadAll(sasPath)\n  selection = "graph selection (prepared before submission)"\nElseIf UCase(mode) = "RUNFILE" Then\n  If project.CodeCollection.Count'
    return VBS.replace(marker,replacement,1).replace('If UCase(mode) <> "RUNFILE" Then','If UCase(mode) <> "RUNFILE" And UCase(mode) <> "RUNGRAPH" Then',1).replace('If UCase(mode) = "RUNFILE" Then\n  On Error Resume Next','If UCase(mode) = "RUNFILE" Or UCase(mode) = "RUNGRAPH" Then\n  On Error Resume Next',1)


def graph_task(task, project, task_root):
    folder=task_root/safe_name(task['task_id']);folder.mkdir(parents=True,exist_ok=True)
    source=task['_graph_source']; template=Path(source['project']); task_project=folder/'connection.egp';shutil.copy2(template,task_project)
    report_progress(folder,'preparing','Preparing '+task['program']+' ('+describe_selection(task)+')')
    chunks=['options iomlogautoflush;']
    if task.get('stop_program_on_error') or any(t.get('stop_program_on_error') for t in task.get('_always_run',[])):
        chunks.append('options errorabend errorcheck=strict;')
    shared=shared_work_code()
    if shared:chunks.append(shared)
    for definition in task.get('_always_run',[]):
        report_progress(folder,'preparing','Appending shared setup: '+definition['program']+' ('+describe_selection(definition)+')')
        chunks.extend(['/* ALWAYS_RUN_ITEM: '+safe_name(definition['task_id'])+' */',read_text(Path(definition['_graph_source']['path']))])
    if shared:chunks.append(shared)
    chunks.extend(['/* SCHEDULED_TARGET_START */',read_text(Path(source['path'])),'/* SCHEDULED_TARGET_END */'])
    submitted=folder/'_submitted.sas';submitted.write_text('\n'.join(chunks)+'\n',encoding='utf-8')
    began=time.time();mode='RUNGRAPH' if source.get('program') else 'RUNFILE'
    rc,console=execute_eg(mode,task_project,submitted,source.get('program') or task['program'],0,0,folder,source.get('tables',False),cancel_file=task.get('_cancel_file'),bridge=graph_bridge() if mode=='RUNGRAPH' else None)
    logs=list((folder/'logs').glob('*.log'));sas_error=any(detect_sas_error(log) for log in logs)
    status='CANCELLED' if rc==130 else 'FAILED' if rc else 'SAS_ERROR' if sas_error else 'SUCCESS'
    evidence=console+'\n'+'\n'.join(read_text(p) for p in logs)
    if any(pattern in evidence.casefold() for pattern in CONNECTION_PATTERNS):status='CONNECTION_LOST'
    return {**task,'status':status,'elapsed':time.time()-began,'task_dir':folder,'message':console.strip().splitlines()[-1] if console.strip() else ''}


def graph_run(args):
    if args.workers is not None and (type(args.workers) is not int or not 1 <= args.workers <= 32):
        raise ValueError('Choose 1 to 32 parallel programs.')
    if args.action=='continue':
        run_dir=Path(args.path).resolve()
        if not (run_dir/'run_summary.csv').is_file():raise ValueError('Choose a previous graph run with saved task results.')
    else:run_dir=prepare_graph_run(args.path,args.workers)
    raw=(run_dir/'graph.json').read_bytes();manifest=json.loads((run_dir/'graph-sources.json').read_text(encoding='utf-8'))
    if hashlib.sha256(raw).hexdigest()!=manifest['graph_hash']:raise ValueError('Graph snapshot changed; create a new run instead.')
    graph,tasks=graph_tasks(run_dir)
    forwarded=argparse.Namespace(workers=args.workers or graph['workers'],no_notify=args.no_notify,cancel_file=getattr(args,'cancel_file',None),prior_elapsed=getattr(args,'prior_elapsed',None))
    if args.action=='continue':forwarded._resume_dir=run_dir
    project=next((Path(t['_graph_source']['project']) for t in tasks if not t['always_run'] and not t['skip']),None)
    with schedule_lock(run_dir):return execute_schedule(forwarded,run_dir,project,tasks)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="pysas.py", description="PySAS — SAS Enterprise Guide command-line utilities")
    p.add_argument("--version", action="version", version=f"PySAS {VERSION}")
    subs = p.add_subparsers(dest="command")
    version = subs.add_parser("version", help="show version and exact script path")
    version.set_defaults(func=lambda a: (print(f"PySAS {VERSION}\nScript: {Path(__file__).resolve()}"), 0)[1])

    bundle = subs.add_parser("bundle", help="pack, verify or unpack SAS code bundles")
    bsub = bundle.add_subparsers(dest="bundle_command", required=True)
    bp = bsub.add_parser("pack"); bp.add_argument("--recursive", action="store_true"); bp.add_argument("--output"); bp.set_defaults(func=bundle_pack)
    bv = bsub.add_parser("verify"); bv.add_argument("--bundle"); bv.set_defaults(func=bundle_verify)
    bu = bsub.add_parser("unpack"); bu.add_argument("--bundle"); bu.add_argument("--no-backup", action="store_true"); bu.set_defaults(func=bundle_unpack)

    for command in (bp, bv, bu):
        command.add_argument("--root", help="code folder to pack, compare or restore (default: folder beside pysas.py)")

    egp = subs.add_parser("egp", help="inspect, extract or repack EGP projects")
    esub = egp.add_subparsers(dest="egp_command", required=True)
    ei = esub.add_parser("inspect"); ei.add_argument("project", nargs="?"); ei.set_defaults(func=egp_inspect)
    ee = esub.add_parser("extract"); ee.add_argument("project", nargs="?"); ee.add_argument("--output"); ee.set_defaults(func=egp_extract)
    ep = esub.add_parser("pack"); ep.add_argument("source"); ep.add_argument("--template"); ep.add_argument("--output"); ep.set_defaults(func=egp_pack)

    runner = subs.add_parser("runner", help="run or watch standalone SAS jobs; _*.sas init files run first")
    rsub = runner.add_subparsers(dest="runner_command", required=True)
    rr = rsub.add_parser("run"); rr.add_argument("program"); rr.add_argument("--template"); rr.add_argument("--lib")
    rr.add_argument("--parameters", help="SAS parameters to prepend after initialization and before the program"); rr.add_argument("--tables", action="store_true"); rr.add_argument("--no-notify", action="store_true"); rr.set_defaults(func=runner_run)
    rw = rsub.add_parser("watch"); rw.add_argument("--template"); rw.add_argument("--lib"); rw.add_argument("--tables", action="store_true")
    rw.add_argument("--workers", type=int, default=2); rw.add_argument("--poll", type=float, default=2.0)
    rw.add_argument("--no-notify", action="store_true"); rw.set_defaults(func=runner_watch)

    rr.add_argument("--init-dir", help="folder containing shared _*.sas initialization files")
    rw.add_argument("--init-dir", help="folder containing shared _*.sas initialization files")
    rw.add_argument("--inbox", help="folder to watch; _*.sas files here initialize jobs instead of being queued")

    graph = subs.add_parser("graph", help="run a JSON dependency graph without an Excel workbook")
    graph.add_argument("action", choices=["run", "continue"]); graph.add_argument("path")
    graph.add_argument("--workers", type=int); graph.add_argument("--no-notify", action="store_true")
    graph.add_argument("--prior-elapsed", type=float, help="Known active seconds for continuing a legacy graph run")
    graph.add_argument("--cancel-file", help=argparse.SUPPRESS); graph.set_defaults(func=graph_run)

    schedule = subs.add_parser("schedule", help="run an Excel dependency schedule")
    schedule.add_argument("action", nargs="?", choices=["run", "continue"], default="run")
    schedule.add_argument("run_dir", nargs="?"); schedule.add_argument("--workbook"); schedule.add_argument("--project")
    schedule.add_argument("--workers", "--max-parallel", type=int, default=None, help="Maximum parallel runs (workbook max_parallel, otherwise 10)"); schedule.add_argument("--no-notify", action="store_true")
    schedule.add_argument("--cancel-file", help=argparse.SUPPRESS)
    schedule.add_argument("--prior-elapsed", type=float, help="Known active seconds for continuing a legacy run")
    schedule.set_defaults(func=lambda a: schedule_continue(a) if a.action == "continue" else schedule_run(a))
    return p


def main(argv: list[str] | None = None) -> int:
    if argv is None: argv = sys.argv[1:]
    if not argv:
        print_home(); return 0
    args = parser().parse_args(argv)
    if not hasattr(args, "func"):
        parser().print_help(); return 2
    try:
        return int(args.func(args) or 0)
    except KeyboardInterrupt:
        print("\nCancelled.", file=sys.stderr); return 130
    except Exception as exc:
        print(f"PySAS error: {exc}", file=sys.stderr); return 1


if __name__ == "__main__":
    raise SystemExit(main())
