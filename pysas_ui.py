"""Local PySAS workbench. Python 3.10+, standard library only."""
from __future__ import annotations
import argparse
import csv
import importlib.util
import io
import json
import mimetypes
import os
import re
import secrets
import shutil
import socketserver
import subprocess
import sys
import threading
import time
import uuid
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlsplit
from pysas import text_encoding
from ui_support import KeepAwake, receive_upload, UPLOAD_LIMIT, launch_app_window

VERSION = "0.4.0-preview.15"
APP_DIR = Path(__file__).resolve().parent
ACTIVE = {"RUNNING", "STOPPING"}


def read_json(path, default=None):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


def read_text(path, limit=400_000, tail=False):
    with path.open("rb") as handle:
        sample = handle.read(min(limit, 4096))
        encoding = text_encoding(sample)
        if encoding == "utf-16":
            encoding = "utf-16-le" if sample.startswith(b"\xff\xfe") else "utf-16-be"
        size = handle.seek(0, 2)
        start = max(0, size - limit) if tail else 0
        if encoding.startswith("utf-16"):
            start -= start % 2
        handle.seek(start)
        raw = handle.read(limit)
    note = "[Showing latest output; download the complete file.]\n" if start else ""
    suffix = "\n[Preview truncated. Download the file for all content.]" if not tail and size > limit else ""
    if encoding == "utf-8-sig":
        import codecs
        if start:
            while raw and raw[0] & 0xc0 == 0x80:
                raw = raw[1:]
        try:
            text = codecs.getincrementaldecoder(encoding)().decode(raw, final=False)
        except UnicodeDecodeError:
            text = raw.decode("cp1252", errors="replace")
    else:
        text = raw.decode(encoding, errors="replace")
    if path.name == "console.txt":
        text = re.sub(r"(?m)^PYSAS_STAGE\|[^|]+\|", "", text)
    return note + text.lstrip("\ufeff") + suffix


def within(root, value):
    candidate = (root / value).resolve()
    if not candidate.is_relative_to(root.resolve()):
        raise ValueError("Choose a path inside the PySAS workspace.")
    return candidate


def atomic_json(path, value):
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


class Workbench:
    def __init__(self, workspace):
        self.root = workspace.resolve()
        self.script = self.root / "pysas.py"
        if not self.script.is_file():
            raise ValueError("The workspace must contain pysas.py.")
        self.storage = self.root / ".pysas-ui"
        self.storage.mkdir(exist_ok=True)
        (self.root / "runner" / "inbox").mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.processes = {}
        self.commands = {}
        self.list_cache = {}
        self.cache_loading = set()
        self.cache_generation = 0
        self.cache_lock = threading.RLock()
        self.awake = KeepAwake()
        self.upload_lock = threading.Lock()
        self.uploading = False
        for path in sorted(self.storage.glob("*.json")):
            item = read_json(path)
            if not isinstance(item, dict) or "id" not in item:
                continue
            if item.get("status") in ACTIVE:
                item["status"] = "UNKNOWN"
                item["message"] = "UI restarted; this process is no longer monitored. Check its output before rerunning."
                for task in item.get("tasks", {}).values():
                    if task.get("status") == "RUNNING":
                        task["status"] = "UNKNOWN"
            self.commands[item["id"]] = item

    def relative(self, path):
        try:
            return str(Path(path).resolve().relative_to(self.root))
        except ValueError:
            return ""

    def save(self, item):
        atomic_json(self.storage / (item["id"] + ".json"), item)

    def folders(self):
        defaults = {"inputs": str(self.root), "init": str(self.root), "inbox": str(self.root / "runner" / "inbox")}
        saved = read_json(self.storage / "folders.json", {})
        return {key: saved.get(key) or value for key, value in defaults.items()}

    def update_folders(self, data):
        with self.lock:
            if self.uploading or any(c["status"] in ACTIVE for c in self.commands.values()):
                raise ValueError("Let active commands finish and stop the watcher before changing folders.")
            settings = self.folders()
            for key in settings:
                if key in data:
                    default = self.root / "runner" / "inbox" if key == "inbox" else self.root
                    raw = str(data[key]).strip()
                    candidate = Path(raw).expanduser() if raw else default
                    if not candidate.is_absolute():
                        candidate = self.root / candidate
                    candidate = candidate.resolve()
                    if not candidate.is_dir():
                        raise ValueError("Create this folder first, then select it: " + str(candidate))
                    settings[key] = str(candidate)
            atomic_json(self.storage / "folders.json", settings)
            return settings

    def input_path(self, value):
        path = (self.root / value).expanduser().resolve()
        roots = [self.root, *(Path(p).resolve() for p in self.folders().values())]
        if not any(path.is_relative_to(root) for root in roots):
            raise ValueError("Select a file inside the workspace or a configured input folder.")
        return path

    def upload(self, query, stream, length):
        destination = query.get("target", ["inputs"])[0]
        folders = self.folders()
        if destination == "bundle":
            folder = self.code_root(query.get("code_root", [""])[0])
        elif destination in folders:
            folder = Path(folders[destination])
            if destination == "inbox":
                folder.mkdir(parents=True, exist_ok=True)
        else:
            raise ValueError("Unknown upload destination.")
        name = query.get("name", [""])[0]
        if destination in {"inbox", "init"} and Path(name).suffix.lower() != ".sas":
            raise ValueError("The inbox and initialization folders accept SAS files only.")
        if destination == "init" and not name.startswith("_"):
            raise ValueError("Shared initialization filenames must start with an underscore.")
        with self.upload_lock:
            self.uploading = True
            try:
                return receive_upload(folder, name, stream, length)
            finally:
                self.uploading = False

    def bundle_settings(self):
        settings = read_json(self.storage / "bundle-paths.json", {})
        return {"paths": settings.get("paths", []), "last": settings.get("last", "")}

    def code_root(self, value):
        root = Path(str(value).strip()).expanduser() if value else self.root
        if not root.is_absolute():
            root = self.root / root
        root = root.resolve()
        if not root.is_dir():
            raise ValueError("Code folder does not exist: " + str(root))
        return root

    def update_bundle_paths(self, data):
        operation = data.get("operation", "save")
        if operation not in {"save", "remove"}:
            raise ValueError("Unknown saved-path action.")
        with self.lock:
            settings = self.bundle_settings()
            if operation == "save":
                path = str(self.code_root(data.get("path")))
                if path not in settings["paths"]:
                    if len(settings["paths"]) >= 20:
                        raise ValueError("You can save up to 20 code folders.")
                    settings["paths"].append(path)
                settings["last"] = path
            else:
                path = str(data.get("path", ""))
                settings["paths"] = [p for p in settings["paths"] if p != path]
                if settings["last"] == path:
                    settings["last"] = ""
            atomic_json(self.storage / "bundle-paths.json", settings)
            return settings

    def arguments(self, data):
        action = data.get("action", "")
        choices = {
            "run": ["runner", "run"], "watch": ["runner", "watch"],
            "schedule": ["schedule"], "continue": ["schedule", "continue"],
            "bundle-pack": ["bundle", "pack"], "bundle-verify": ["bundle", "verify"],
            "bundle-unpack": ["bundle", "unpack"], "egp-inspect": ["egp", "inspect"],
            "egp-extract": ["egp", "extract"], "egp-pack": ["egp", "pack"],
        }
        if action not in choices:
            raise ValueError("Unknown command.")
        if action in {"run", "watch", "schedule", "continue"} and os.name != "nt":
            raise ValueError("SAS execution requires Windows and SAS Enterprise Guide. File tools and history work here.")
        data = dict(data)
        folders = self.folders()
        project_field = "template" if action in {"run", "watch", "egp-pack"} else "project"
        if action in {"run", "watch", "schedule", "egp-pack", "egp-inspect", "egp-extract"} and not data.get(project_field) and Path(folders["inputs"]) != self.root:
            projects = sorted(Path(folders["inputs"]).glob("*.egp"))
            if len(projects) != 1:
                raise ValueError("Select an EGP project from the input folder.")
            data[project_field] = str(projects[0])
        args = choices[action].copy()
        if action in {"run", "watch"}:
            init_dir = Path(folders["init"])
            if not init_dir.is_dir():
                raise ValueError("The configured initialization folder is unavailable.")
            args.extend(["--init-dir", str(init_dir)])
        if action == "watch":
            args.extend(["--inbox", folders["inbox"]])
        file_root = self.code_root(data.get("code_root")) if action.startswith("bundle-") else self.root
        if action.startswith("bundle-") and data.get("code_root"):
            args.extend(["--root", str(file_root)])
        required = {"run": "program", "continue": "run_dir", "egp-pack": "source"}
        if action in required:
            value = str(data.get(required[action], "")).strip()
            if not value:
                raise ValueError("Select " + required[action].replace("_", " ") + ".")
            path = self.input_path(value)
            if not path.exists():
                raise ValueError("The selected path no longer exists.")
            args.append(str(path))
        if action in {"egp-inspect", "egp-extract"} and data.get("project"):
            args.append(str(self.input_path(data["project"])))
        allowed = {
            "run": ["template", "lib"], "watch": ["template", "lib"],
            "schedule": ["workbook", "project"], "continue": [],
            "bundle-pack": ["output"], "bundle-verify": ["bundle"],
            "bundle-unpack": ["bundle"], "egp-inspect": [],
            "egp-extract": ["output"], "egp-pack": ["template", "output"],
        }
        for field in allowed[action]:
            if data.get(field):
                path = within(file_root, str(data[field]).strip()) if action.startswith("bundle-") or field == "output" else self.input_path(str(data[field]).strip())
                if field != "output" and not path.is_file():
                    raise ValueError("File not found: " + str(data[field]))
                if field == "output":
                    protected = {self.script, APP_DIR / "pysas_ui.py", APP_DIR / "ui_worker.py"}
                    if path in protected or path == self.storage or self.storage in path.parents or path.suffix.lower() in {".py", ".js", ".css", ".bat", ".cmd", ".exe", ".dll", ".ps1"}:
                        raise ValueError("Choose an output path outside application files.")
                args.extend(["--" + field, str(path)])
        if action in {"watch", "schedule", "continue"}:
            workers = int(data.get("workers") or (2 if action == "watch" else 10))
            if not 1 <= workers <= 32:
                raise ValueError("Workers must be between 1 and 32.")
            args.extend(["--workers", str(workers)])
        if action == "watch":
            poll = float(data.get("poll") or 2)
            if not 0.5 <= poll <= 60:
                raise ValueError("Polling interval must be between 0.5 and 60 seconds.")
            args.extend(["--poll", str(poll)])
        if action in {"run", "watch"} and data.get("tables"):
            args.append("--tables")
        if action in {"run", "watch", "schedule", "continue"}:
            args.append("--no-notify")
        if action == "bundle-pack" and data.get("recursive"):
            args.append("--recursive")
        return action, args

    def launch(self, data):
        action, args = self.arguments(data)
        script = self.script
        with self.lock:
            if getattr(self, "closing", False):
                raise ValueError("PySAS is closing and finishing its active jobs.")
            if action == "watch" and any(c["action"] == "watch" and c["status"] in ACTIVE for c in self.commands.values()):
                raise ValueError("The watcher is already running in this workbench.")
            identifier = uuid.uuid4().hex
            item = {"id": identifier, "action": action, "args": args,
                    "name": data.get("program") or data.get("workbook") or data.get("run_dir") or action.replace("-", " "),
                    "code_root": str(self.code_root(data.get("code_root"))) if action.startswith("bundle-") else None,
                    "started": time.time(), "finished": None, "status": "RUNNING", "tasks": {}, "message": "",
                    "can_cancel_file": True}
            if action == "bundle-pack":
                root = self.code_root(data.get("code_root"))
                item["artifact"] = str(within(root, str(data.get("output") or "codebase.sasbundle.txt")))
            env = os.environ.copy()
            env.update(PYTHONUNBUFFERED="1", PYTHONIOENCODING="utf-8")
            env.pop("PYSAS_LIVE_LOG", None)
            isolation = {"creationflags": subprocess.CREATE_NEW_PROCESS_GROUP | subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {"start_new_session": True}
            if os.name == "nt" and action in {"run", "watch", "schedule", "continue"}:
                # Keep the same console semantics as terminal Python, without a visible window.
                startup = subprocess.STARTUPINFO()
                startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
                startup.wShowWindow = subprocess.SW_HIDE
                isolation = {"creationflags": subprocess.CREATE_NEW_CONSOLE, "startupinfo": startup}
                env["PYSAS_UI_CONSOLE"] = "1"
            else:
                env.pop("PYSAS_UI_CONSOLE", None)
            control_path = self.storage / (identifier + ".stop")
            env["PYSAS_UI_CONTROL_FILE"] = str(control_path)
            item["control_file"] = str(control_path)
            python = str(Path(sys.executable).with_name("python.exe")) if os.name == "nt" else sys.executable
            event_path = self.storage / (identifier + ".events.jsonl")
            event_path.touch()
            env["PYSAS_UI_EVENTS"] = str(event_path)
            env.pop("PYSAS_UI_LEGACY_ROOT", None)
            # Neither normal output nor events depend on the HTTP/history reader.
            with (self.storage / (identifier + ".txt")).open("wb", buffering=0) as output:
                process = subprocess.Popen([python, "-u", str(APP_DIR / "ui_worker.py"), str(script), *args],
                                           cwd=self.root, stdin=None, stdout=output,
                                           stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", env=env, **isolation)
            self.commands[identifier] = item
            self.processes[identifier] = process
            try:
                self.save(item)
            except OSError as exc:
                item["message"] = "UI history could not be saved: " + str(exc)
            threading.Thread(target=self.monitor_worker, args=(identifier, process, event_path), daemon=True).start()
            if action.startswith("bundle-"):
                try:
                    settings = self.bundle_settings()
                    settings["last"] = item["code_root"]
                    atomic_json(self.storage / "bundle-paths.json", settings)
                except OSError as exc:
                    item["message"] = "Bundle folder preference could not be saved: " + str(exc)
            return {"id": identifier}

    def event(self, item, event):
        kind = event.get("event")
        if kind == "worker":
            item["runtime"] = {k: v for k, v in event.items() if k != "event"}
        elif kind == "plan":
            for task in event["tasks"]:
                key = task["task_id"]
                if key not in item["tasks"]:
                    item["tasks"][key] = {"key": key, "name": task["program"], "kind": "task",
                        "status": "SKIPPED_SUCCESS" if task.get("skip") else ("ALWAYS_RUN_DEFINITION" if task.get("always_run") else "PENDING"),
                        "depends_on": task.get("depends_on", []), "started": None, "elapsed": 0,
                        "section": task.get("section", ""), "row_start": task.get("row_start"), "row_end": task.get("row_end")}
        elif kind in {"start", "finish"}:
            key = event["key"]
            task = item["tasks"].setdefault(key, {"key": key})
            task["can_cancel_file"] = item.get("can_cancel_file", True)
            task.update({k: v for k, v in event.items() if k not in {"event", "time", "path"}})
            if event.get("path"):
                task["path"] = self.relative(event["path"])
            if kind == "start":
                task.update(status="RUNNING", started=event["time"])
            else:
                task["finished"] = event["time"]
                task["elapsed"] = event.get("elapsed", max(0, event["time"] - (task.get("started") or event["time"])))
        elif kind == "progress":
            task = item["tasks"].get(event["key"])
            if task is not None:
                task.update(phase=event["phase"], progress=event["progress"], phase_started=event["time"])
                if event.get("path"):
                    task["path"] = self.relative(event["path"])
        elif kind == "location":
            task = item["tasks"].get(event["key"])
            if task is not None:
                task["path"] = self.relative(event["path"])
        elif kind == "summary":
            item["path"] = self.relative(event["path"])
            for result in event["tasks"]:
                task = item["tasks"].setdefault(result["task_id"], {"key": result["task_id"], "name": result["program"]})
                task.update({field: result[field] for field in ("section", "row_start", "row_end") if field in result})
                task.update(status=result["status"], elapsed=result.get("elapsed", 0), message=result.get("message", ""))
                if result.get("task_dir"):
                    task["path"] = self.relative(result["task_dir"])
        self.save(item)

    def monitor_worker(self, identifier, process, event_path):
        """Tail complete events; process exit, never inherited-file EOF, ends a run."""
        pending = ""
        try:
            with event_path.open(encoding="utf-8", errors="replace") as events:
                while True:
                    exited = process.poll() is not None
                    chunk = events.read(65536)
                    pending += chunk
                    lines = pending.split("\n")
                    pending = lines.pop()
                    for line in lines:
                        try:
                            event = json.loads(line)
                            with self.lock:
                                self.event(self.commands[identifier], event)
                        except (ValueError, KeyError, OSError) as exc:
                            # A history write failure cannot terminate observation or SAS.
                            with self.lock:
                                self.commands[identifier]["message"] = "UI history update failed: " + str(exc)
                    if exited and not chunk:
                        break
                    if not chunk:
                        time.sleep(.1)
        except OSError as exc:
            with self.lock:
                self.commands[identifier]["message"] = "UI event file could not be read: " + str(exc)
        finally:
            self.complete_worker(identifier, process)

    def complete_worker(self, identifier, process):
        rc = process.wait()
        with self.lock:
            item = self.commands[identifier]
            item["finished"] = time.time()
            item["elapsed"] = item["finished"] - item["started"]
            item["exit_code"] = rc
            if rc == 0 and item["action"] == "bundle-pack" and item.get("artifact"):
                source = within(Path(item["code_root"]), item["artifact"])
                target = self.storage / "artifacts" / identifier / source.name
                try:
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(source, target)
                    item["download"] = str(target.relative_to(self.storage / "artifacts"))
                except OSError as exc:
                    item["message"] = "Bundle created, but download copy failed: " + str(exc)
            item["status"] = "STOPPED" if item["status"] == "STOPPING" and rc in {0, 130} else ("SUCCESS" if rc == 0 else "FAILED")
            for task in item["tasks"].values():
                if task.get("status") in {"RUNNING", "PENDING", "CANCELLING"}:
                    task["status"] = "UNKNOWN"
                    task["message"] = "No final task status was received; check the console."
            try:
                self.save(item)
            except OSError as exc:
                item["message"] = "Run ended, but history could not be saved: " + str(exc)
            finally:
                self.processes.pop(identifier, None)
        if process.stdin is not None:
            process.stdin.close()

    def stop(self, identifier):
        with self.lock:
            item = self.commands.get(identifier)
            process = self.processes.get(identifier)
            if not item or not process or item["action"] != "watch":
                raise ValueError("Only an active watcher can be stopped here.")
            if item["status"] == "STOPPING":
                return {"ok": True}
            Path(item["control_file"]).write_text("stop", encoding="utf-8")
            item["status"] = "STOPPING"
            self.save(item)
            return {"ok": True}

    def cancel_file(self, identifier, key):
        with self.lock:
            item = self.commands.get(identifier)
            task = item.get("tasks", {}).get(key) if item else None
            if identifier not in self.processes or not task or task.get("status") not in {"RUNNING", "CANCELLING"}:
                raise ValueError("This file is no longer running.")
            if item.get("can_cancel_file") is False:
                raise ValueError("This historical run does not support per-file cancellation.")
            if not task.get("path"):
                raise ValueError("The file is still starting. Try again in a moment.")
            run_dir = within(self.root, task["path"])
            if not run_dir.is_dir():
                raise ValueError("The run folder is not available yet.")
            (run_dir / "_cancel.request").write_text("Requested by user", encoding="utf-8")
            task["status"] = "CANCELLING"
            task["message"] = "Stopping this file's automation process…"
            self.save(item)
            return {"ok": True}

    def cached_listing(self, name, loader, ttl):
        # Slow/network directory scans must never hold up live command status.
        with self.cache_lock:
            cached = self.list_cache.get(name)
            generation = self.cache_generation
            key = (name, generation)
            if (not cached or time.monotonic() - cached[0] >= ttl) and key not in self.cache_loading:
                self.cache_loading.add(key)
                def update():
                    try:
                        value = loader()
                        with self.cache_lock:
                            if generation == self.cache_generation:
                                self.list_cache[name] = (time.monotonic(), value)
                    except OSError:
                        pass  # Retain the previous listing if a network folder is unavailable.
                    finally:
                        with self.cache_lock:
                            self.cache_loading.discard(key)
                threading.Thread(target=update, daemon=True).start()
            return cached[1] if cached else []

    def invalidate_lists(self):
        with self.cache_lock:
            self.cache_generation += 1
            self.list_cache.clear()

    def inventory(self):
        paths = set()
        roots = {self.root, Path(self.folders()["inputs"]), Path(self.folders()["init"])}
        for root in sorted(roots):
            for base, dirs, names in os.walk(root):
                dirs[:] = sorted(d for d in dirs if not d.startswith(".") and d not in {"runner", "runs", "__pycache__", "_codebase_backups", "node_modules", "release", "ui", "tests"})
                for name in sorted(names):
                    path = Path(base) / name
                    if path.suffix.lower() in {".sas", ".egp", ".xlsx", ".txt", ".csv"} and path.resolve().is_relative_to(root.resolve()):
                        paths.add(self.relative(path) or str(path.resolve()))
                    if len(paths) >= 3000:
                        return sorted(paths)
        return sorted(paths)

    def history(self):
        records = []
        for root, kind in [(self.root / "runner" / "runs", "file"), (self.root / "runs", "schedule")]:
            if not root.exists():
                continue
            for path in root.iterdir():
                if not path.is_dir() or not path.resolve().is_relative_to(self.root):
                    continue
                row = {"path": self.relative(path), "name": path.name, "kind": kind,
                       "status": "UNKNOWN", "elapsed": None, "started": path.stat().st_mtime}
                status = path / "status.txt"
                if status.is_file():
                    values = dict(line.split("=", 1) for line in read_text(status).splitlines() if "=" in line)
                    row.update(name=values.get("source", path.name), status=values.get("status", "UNKNOWN"))
                    try:
                        row["elapsed"] = float(values["elapsed_seconds"])
                        row["started"] = datetime.fromisoformat(values["started"]).timestamp()
                    except (KeyError, ValueError):
                        pass
                summary = path / "run_summary.csv"
                if summary.is_file():
                    tasks = list(csv.DictReader(io.StringIO(read_text(summary))))
                    row["tasks"] = len(tasks)
                    row["status"] = "SUCCESS" if all(t["status"] in {"SUCCESS", "SKIPPED_SUCCESS", "SKIPPED_PREVIOUS", "ALWAYS_RUN_DEFINITION"} for t in tasks) else "FAILED"
                records.append(row)
        return sorted(records, key=lambda x: x["started"], reverse=True)[:500]

    def state(self, refresh=False):
        if refresh:
            self.invalidate_lists()
        with self.lock:
            commands = json.loads(json.dumps(sorted(self.commands.values(), key=lambda x: x["started"], reverse=True)[:200]))
        history = [dict(row) for row in self.cached_listing("history", self.history, 15)]
        paths = {h["path"]: h for h in history}
        for command in commands:
            if command.get("path") in paths:
                paths[command["path"]].update(elapsed=command.get("elapsed"), started=command["started"])
            for task in command.get("tasks", {}).values():
                if task.get("path") in paths:
                    paths[task["path"]].update(status=task["status"], elapsed=task.get("elapsed"), started=task.get("started") or paths[task["path"]]["started"])
        inbox = Path(self.folders()["inbox"])
        queued = self.cached_listing("queue", lambda: sorted(p.name for p in inbox.iterdir() if p.is_file() and p.suffix.casefold() == ".sas" and not p.name.startswith("_")) if inbox.exists() else [], 2)
        return {"version": VERSION, "workspace": str(self.root), "windows": os.name == "nt",
                "openpyxl": importlib.util.find_spec("openpyxl") is not None,
                "files": self.cached_listing("inventory", self.inventory, 30), "commands": commands[:200], "history": history,
                "folders": self.folders(), "awake": self.awake.status(),
                "initialization_files": self.cached_listing("init", lambda: sorted({str(p.resolve()) for folder in {Path(self.folders()["init"]), inbox} for p in folder.glob("_*") if p.is_file() and p.suffix.casefold() == ".sas"}), 5),
                "queued": queued, "bundle_settings": self.bundle_settings(), "now": time.time()}

    def details(self, relative):
        path = within(self.root, relative)
        if not path.is_dir() or not relative:
            raise ValueError("Choose a run folder.")
        files = []
        for file in sorted(path.rglob("*")):
            if file.is_file() and file.resolve().is_relative_to(self.root):
                files.append({"path": self.relative(file), "name": str(file.relative_to(path)), "size": file.stat().st_size})
            if len(files) >= 1000:
                break
        summary = path / "run_summary.csv"
        tasks = list(csv.DictReader(io.StringIO(read_text(summary)))) if summary.is_file() else []
        with self.lock:
            for command in self.commands.values():
                if command.get("path") == self.relative(path):
                    for task in tasks:
                        saved = command.get("tasks", {}).get(task["task_id"], {})
                        task.update({field: saved[field] for field in ("section", "row_start", "row_end") if field in saved})
                    break
        return {"files": files, "tasks": tasks}

    def preview(self, relative):
        path = self.input_path(relative)
        if path.suffix.lower() == ".xlsx":
            try:
                import openpyxl
            except ImportError:
                raise ValueError("Install openpyxl to preview Excel workbooks; you can still download the file.")
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                sheets = []
                for sheet in workbook.worksheets[:10]:
                    rows = []
                    for row in sheet.iter_rows(max_row=101, max_col=min(sheet.max_column or 1, 30), values_only=True):
                        rows.append([str(v) if v is not None else "" for v in row])
                    sheets.append({"name": sheet.title, "rows": rows})
                return {"sheets": sheets, "note": "Preview: up to 100 data rows, 30 columns, and 10 sheets."}
            finally:
                workbook.close()
        if path.suffix.lower() not in {".sas", ".log", ".txt", ".csv", ".tsv", ".json", ".html", ".htm", ".xml"}:
            raise ValueError("Download this file to open it in its application.")
        return {"text": read_text(path, limit=60000 if path.suffix.lower() == ".log" or path.name == "console.txt" else 400000, tail=path.suffix.lower() == ".log" or path.name == "console.txt"),
                "modified": path.stat().st_mtime, "size": path.stat().st_size}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def respond(self, value, status=200, content_type="application/json; charset=utf-8"):
        body = json.dumps(value, default=str).encode() if isinstance(value, (dict, list)) else value
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'none'; form-action 'self'")
        self.end_headers()
        self.wfile.write(body)

    def authorized(self):
        host = self.headers.get("Host", "")
        expected = "127.0.0.1:" + str(self.server.server_port)
        if host != expected:
            self.respond({"error": "Invalid host."}, 403)
            return False
        origin = self.headers.get("Origin")
        if origin and origin != "http://" + expected:
            self.respond({"error": "Invalid origin."}, 403)
            return False
        return True

    def do_GET(self):
        if not self.authorized():
            return
        url = urlsplit(self.path)
        query = parse_qs(url.query)
        value = lambda key: query.get(key, [""])[0]
        try:
            if url.path == "/api/state":
                self.respond({**self.server.app.state(refresh=value("refresh") == "1"), "token": self.server.token})
            elif url.path == "/api/details":
                self.respond(self.server.app.details(value("path")))
            elif url.path == "/api/preview":
                self.respond(self.server.app.preview(value("path")))
            elif url.path == "/api/console":
                identifier = value("id")
                if identifier not in self.server.app.commands:
                    raise ValueError("Command not found.")
                path = self.server.app.storage / (identifier + ".txt")
                if path.exists():
                    with path.open("rb") as f:
                        f.seek(max(0, path.stat().st_size - 100_000))
                        console = f.read().decode("utf-8", errors="replace")
                else:
                    console = "Waiting for output…"
                with self.server.app.lock:
                    tasks = list(self.server.app.commands[identifier].get("tasks", {}).values())[-8:]
                for task in tasks:
                    if task.get("path"):
                        stage = within(self.server.app.root, task["path"]) / "console.txt"
                        if stage.is_file():
                            console += "\n\n--- " + str(task.get("name", task.get("key", "File"))) + " ---\n" + read_text(stage, limit=8000, tail=True)
                self.respond({"text": console})
            elif url.path == "/api/download":
                if value("command"):
                    item = self.server.app.commands.get(value("command"), {})
                    if item.get("action") != "bundle-pack" or item.get("status") != "SUCCESS" or not item.get("download"):
                        raise ValueError("No completed bundle is available for this command.")
                    path = within(self.server.app.storage / "artifacts", item["download"])
                else:
                    path = self.server.app.input_path(value("path"))
                if not path.is_file():
                    raise ValueError("File not found.")
                self.send_response(200)
                self.send_header("Content-Type", "application/octet-stream")
                self.send_header("Content-Disposition", "attachment; filename*=UTF-8''" + quote(path.name))
                self.send_header("Content-Length", str(path.stat().st_size))
                self.send_header("X-Content-Type-Options", "nosniff")
                self.end_headers()
                with path.open("rb") as f:
                    while chunk := f.read(256 * 1024):
                        self.wfile.write(chunk)
            else:
                name = {"/": "index.html", "/app.js": "app.js", "/style.css": "style.css", "/icon.svg": "icon.svg"}.get(url.path)
                if not name:
                    self.respond({"error": "Not found."}, 404)
                    return
                path = APP_DIR / "ui" / name
                self.respond(path.read_bytes(), content_type=mimetypes.guess_type(name)[0] or "text/plain")
        except (ValueError, OSError, KeyError) as exc:
            self.respond({"error": str(exc)}, 400)

    def do_POST(self):
        if not self.authorized():
            return
        if not secrets.compare_digest(self.headers.get("X-PySAS-Token", ""), self.server.token):
            self.respond({"error": "Refresh the workbench before trying again."}, 403)
            return
        try:
            self.server.app.invalidate_lists()
            length = int(self.headers.get("Content-Length", "0"))
            url = urlsplit(self.path)
            if url.path == "/api/upload":
                if not 0 <= length <= UPLOAD_LIMIT:
                    raise ValueError("Each file must be 512 MB or smaller.")
                self.connection.settimeout(60)
                self.respond(self.server.app.upload(parse_qs(url.query), self.rfile, length))
                return
            if length < 0 or length > 64_000:
                raise ValueError("Request is too large.")
            data = json.loads(self.rfile.read(length))
            if not isinstance(data, dict):
                raise ValueError("Expected an object.")
            if self.path == "/api/quit":
                self.respond({"message": "Closing PySAS after active jobs finish."})
                threading.Thread(target=self.server.request_close, daemon=True).start()
            elif self.path == "/api/launch":
                self.respond(self.server.app.launch(data))
            elif self.path == "/api/folders":
                self.respond(self.server.app.update_folders(data))
            elif self.path == "/api/awake":
                self.respond(self.server.app.awake.set(data.get("enabled")))
            elif self.path == "/api/bundle-paths":
                self.respond(self.server.app.update_bundle_paths(data))
            elif self.path == "/api/cancel-file":
                self.respond(self.server.app.cancel_file(data.get("id"), data.get("key")))
            elif self.path == "/api/stop":
                self.respond(self.server.app.stop(data.get("id")))
            else:
                self.respond({"error": "Not found."}, 404)
        except (ValueError, TypeError, OSError) as exc:
            self.respond({"error": str(exc)}, 400)


class LocalServer(ThreadingHTTPServer):
    def request_close(self):
        with self.app.lock:
            if getattr(self.app, "closing", False):
                return
            self.app.closing = True
            watchers = [key for key in self.app.processes if self.app.commands[key]["action"] == "watch"]
        for key in watchers:
            try:
                self.app.stop(key)
            except (OSError, ValueError):
                pass
        while self.app.processes:
            time.sleep(.2)
        self.shutdown()

    def server_bind(self):
        # Avoid reverse DNS lookups: this service is explicitly loopback-only.
        socketserver.TCPServer.server_bind(self)
        self.server_name = "127.0.0.1"
        self.server_port = self.server_address[1]

    def server_close(self):
        super().server_close()
        if hasattr(self, "app"):
            self.app.awake.close()
        handle = getattr(self, "workspace_lock", None)
        if handle is not None:
            handle.close()
            self.workspace_lock = None


def make_server(workspace, port=0):
    storage = workspace.resolve() / ".pysas-ui"
    storage.mkdir(exist_ok=True)
    handle = (storage / "server.lock").open("a+b")
    try:
        if os.name == "nt":
            import msvcrt
            if handle.seek(0, 2) == 0:
                handle.write(b"0")
                handle.flush()
            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        handle.close()
        raise ValueError("A workbench is already running for this workspace. Use its browser window or stop its launcher first.")
    try:
        app = Workbench(workspace)
        server = LocalServer(("127.0.0.1", port), Handler)
    except Exception:
        handle.close()
        raise
    server.workspace_lock = handle
    server.app = app
    server.token = secrets.token_urlsafe(32)
    return server


def main():
    parser = argparse.ArgumentParser(description="Start the local PySAS workbench")
    parser.add_argument("--workspace", type=Path, default=APP_DIR)
    parser.add_argument("--port", type=int, default=0)
    window_options = parser.add_mutually_exclusive_group()
    window_options.add_argument("--no-browser", action="store_true", help="start the server without opening a window")
    window_options.add_argument("--browser", action="store_true", help="open a normal browser tab instead of the Windows app window")
    parser.add_argument("--ready-file", type=Path, help=argparse.SUPPRESS)
    args = parser.parse_args()
    if sys.stdout is None or sys.stderr is None:
        storage = args.workspace.resolve() / ".pysas-ui"
        storage.mkdir(parents=True, exist_ok=True)
        output = (storage / "launcher.log").open("a", encoding="utf-8", buffering=1)
        sys.stdout = sys.stderr = output
    server = make_server(args.workspace, args.port)
    url = f"http://127.0.0.1:{server.server_port}"
    print(f"PySAS Workbench {VERSION}\nWorkspace: {args.workspace.resolve()}\n{url}", flush=True)
    monitor_stop = threading.Event()
    window_process = None
    def ready():
        if args.ready_file:
            temp = args.ready_file.with_suffix(".tmp")
            temp.write_text(url, encoding="utf-8")
            temp.replace(args.ready_file)
    def watch_window(process):
        from windows_app import process_windows, brand_window
        opened = False
        deadline = time.monotonic() + 45
        missing_since = None
        branded = {}
        while not monitor_stop.wait(.5):
            windows = process_windows(process.pid)
            if windows:
                missing_since = None
                for hwnd in windows:
                    count, last = branded.get(hwnd, (0, 0))
                    if count >= 3 or time.monotonic() - last < 3:
                        continue
                    try:
                        brand_window(hwnd, args.workspace, APP_DIR / "ui" / "icon.ico")
                        branded[hwnd] = (count + 1, time.monotonic())
                    except OSError as exc:
                        print(f"Taskbar identity: {exc}", flush=True)
                if not opened:
                    opened = True
                    ready()
            elif opened:
                missing_since = missing_since or time.monotonic()
                if time.monotonic() - missing_since > 2:
                    server.request_close()
                    return
            elif time.monotonic() > deadline or process.poll() is not None:
                print("The app window could not be opened. Use START_PYSAS_BROWSER.bat or inspect launcher.log.", flush=True)
                server.request_close()
                return
    try:
        if not args.no_browser:
            if os.name == "nt" and not args.browser:
                window_process = launch_app_window(url, server.app.storage / "app-profile")
                if window_process is None:
                    raise RuntimeError("Install Microsoft Edge or Google Chrome, or use START_PYSAS_BROWSER.bat.")
                threading.Thread(target=watch_window, args=(window_process,), daemon=True).start()
            else:
                if not webbrowser.open(url):
                    raise RuntimeError("Could not open the browser. Use START_PYSAS_CONSOLE.bat --no-browser.")
                ready()
        else:
            ready()
        server.serve_forever(poll_interval=0.5)
    except KeyboardInterrupt:
        # shutdown() must run outside the serve_forever thread.
        threading.Thread(target=server.request_close, daemon=True).start()
        while server.app.processes:
            time.sleep(.2)
    finally:
        monitor_stop.set()
        if window_process is not None:
            from windows_app import close_windows
            close_windows(window_process.pid)
            try:
                window_process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                print("Browser is still finishing shutdown.", flush=True)
        server.server_close()


if __name__ == "__main__":
    main()
