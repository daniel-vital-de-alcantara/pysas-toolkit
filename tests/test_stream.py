"""Run stream composition, portability, engine boundary and Workbench integration."""
import base64
import copy
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from types import SimpleNamespace
import unittest
from unittest.mock import patch, Mock
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import engine
import pysas_ui as ui
import ui_stream as stream
import stream_worker
from ui_data import duration_history, remaining_work
import ui_storage


class StreamTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name).resolve()
        shutil.copy2(ROOT / 'pysas.py', self.root / 'pysas.py')
        self.app = ui.Workbench(self.root)
        self.egp = self.root / 'connection.egp'; self.egp.write_text('{}')
        self.source = self.root / 'extract.sas'
        self.source.write_text('%let ids=10;\n* (please do not delete) section_start: Realised;\n%put café &ids.;\n* (please do not delete) section_end: Realised;\n%put end;', encoding='utf-8')
        self.doc = dict(format=stream.FORMAT, name='Extract cases', egp=stream.describe(self.app, self.egp), server_program='',
                        rows=[self.row('first')], export_html=True, export_tables=False, consolidate_tables=False)

    def tearDown(self):
        self.app.awake.close(); self.temp.cleanup()

    def row(self, identifier, **values):
        return dict(id=identifier, file=stream.describe(self.app, self.source), mode='whole', row_start=None, row_end=None, section='', **values)

    def plan(self, **values):
        return stream.compose(self.app, {**self.doc, **values})

    def test_order_line_bounds_sections_and_no_hidden_setup(self):
        (self.root / '_hidden.sas').write_text('%put MUST_NOT_RUN;')
        first=self.row('one'); first.update(mode='lines', row_end=1)
        section=self.row('two'); section.update(mode='section', section='realised', row_start=999)
        tail=self.row('three'); tail.update(mode='lines', row_start=5)
        plan=self.plan(rows=[first,section,tail])
        source=plan['source']
        self.assertLess(source.index('%let ids'),source.index('%put café')); self.assertLess(source.index('%put café'),source.index('%put end'))
        self.assertNotIn('MUST_NOT_RUN',source); self.assertNotIn('section_start:',source)
        self.assertEqual([s['scope'] for s in plan['selections']],['Lines 1–1','Section realised','Lines 5–end'])
        self.assertEqual(len(plan['selections']),3)

    def test_one_line_range_is_inclusive_and_whole_file_is_complete(self):
        row=self.row('one'); row.update(mode='lines',row_start=3,row_end=3)
        selected=self.plan(rows=[row])['source']
        self.assertIn('%put café', selected); self.assertNotIn('%let ids',selected); self.assertNotIn('%put end',selected)
        self.assertIn(self.source.read_text(encoding='utf-8'),self.plan()['source'])

    def test_invalid_or_empty_selections_do_not_create_runs(self):
        cases=[[],[None]]
        for values in ({'mode':'lines','row_start':6},{'mode':'lines','row_start':3,'row_end':2},
                       {'mode':'section','section':'missing'},{'mode':'section','section':''},{'row_start':0},{'row_start':True}):
            row=self.row('bad');row.update(values);cases.append([row])
        for rows in cases:
            with self.subTest(rows=rows), self.assertRaises((ValueError,TypeError)):
                stream.prepare(self.app,{**self.doc,'rows':rows})
        self.assertEqual(list((self.app.tasks/'files').glob('*')),[])

    def test_changed_missing_and_disallowed_assets_fail_before_submission(self):
        original=self.source.read_bytes();self.source.write_text('changed;')
        with self.assertRaisesRegex(ValueError,'changed'):self.plan()
        self.source.write_bytes(original);self.source.unlink()
        with self.assertRaisesRegex(ValueError,'missing'):self.plan()
        with self.assertRaises(ValueError):stream.describe(self.app,self.root.parent/'outside.sas')

    def test_legacy_windows_and_utf16_sources_compile_without_mojibake(self):
        for encoding in ('utf-16','utf-16-le','cp1252'):
            self.source.write_bytes('%put preço €;'.encode(encoding));self.doc['rows'][0]['file']=stream.describe(self.app,self.source)
            self.assertIn('%put preço €;',self.plan()['source'])

    def test_portable_json_restores_files_order_choices_and_blank_rows_in_other_installation(self):
        self.doc['rows'][0].update(mode='section',section='Realised',row_start=5,row_end=6)
        self.doc['rows'].append(dict(id='empty',file=None,mode='lines',section='Later',row_start=None,row_end=50))
        self.doc.update(export_html=False,export_tables=True,consolidate_tables=True,server_program='Connection seed')
        exported=stream.artifact(self.app,self.doc,'json'); raw=self.app.resolve(exported['path']).read_bytes()
        second=self.root/'other-install';second.mkdir();shutil.copy2(ROOT/'pysas.py',second/'pysas.py');other=ui.Workbench(second)
        self.source.unlink();self.egp.unlink()
        restored=stream.import_document(other,raw)
        self.assertEqual(restored['rows'][0]['section'],'Realised')
        for key,value in self.doc['rows'][1].items():self.assertEqual(restored['rows'][1][key],value)
        self.assertEqual(restored['rows'][1]['sections'],['Later'])
        self.assertEqual(restored['rows'][0]['row_start'],5);self.assertTrue(restored['consolidate_tables']);self.assertFalse(restored['export_html'])
        self.assertIn('Realised',restored['rows'][0]['file']['sections'])
        self.assertTrue(other.resolve(restored['egp']['path']).is_file());self.assertEqual(other.resolve(restored['rows'][0]['file']['path']).name[-4:],'.sas')
        restored['rows'].pop();self.assertIn('%put café',stream.compose(other,restored,require_egp=True)['source'])

    def test_invalid_import_is_atomic_and_rejects_tampered_embedded_code(self):
        exported=stream.artifact(self.app,self.doc,'json');raw=self.app.resolve(exported['path']).read_bytes();data=json.loads(raw)
        data['assets'][self.doc['rows'][0]['file']['sha256']]=base64.b64encode(b'changed').decode()
        with self.assertRaisesRegex(ValueError,'checksum'):stream.import_document(self.app,json.dumps(data).encode())
        self.assertFalse((self.app.tasks/'stream-sources').exists())
        data=json.loads(raw);data['rows'].append(copy.deepcopy(data['rows'][0]))
        with self.assertRaisesRegex(ValueError,'unique'):stream.import_document(self.app,json.dumps(data).encode())

    def test_uploads_with_same_name_are_distinct_and_keep_original_display_name(self):
        first=self.app.upload({'target':['stream'],'name':['extract.sas']},io.BytesIO(b'%put one;'),9)
        second=self.app.upload({'target':['stream'],'name':['extract.sas']},io.BytesIO(b'%put two;'),9)
        self.assertNotEqual(first['path'],second['path']);self.assertEqual(first['name'],'extract.sas')
        self.assertEqual(self.app.resolve(first['path']).read_bytes(),b'%put one;')

    def test_source_export_is_the_exact_frozen_source_given_to_engine(self):
        exported=stream.artifact(self.app,self.doc,'source')
        folder,_,_=stream.prepare(self.app,self.doc)
        self.assertEqual(self.app.resolve(exported['path']).read_bytes(),(folder/'source.sas').read_bytes())
        original=self.source.read_bytes();self.source.write_text('changed after preparing;')
        doc=json.loads((folder/'stream-plan.json').read_text())['document']
        self.assertEqual((folder/doc['rows'][0]['file']['path']).read_bytes(),original)
        self.assertTrue((folder/'stream.json').exists())

    def run_fake_engine(self, options=None):
        self.egp.write_text(json.dumps(options or {}));self.doc['egp']=stream.describe(self.app,self.egp)
        folder,_,_=stream.prepare(self.app,self.doc)
        events=[]
        def launch(request,output):
            isolation={'creationflags':subprocess.CREATE_NEW_PROCESS_GROUP} if os.name=='nt' else {'start_new_session':True}
            return subprocess.Popen([sys.executable,str(ROOT/'tests/fixtures/standalone_engine_worker.py'),str(request)],stdout=output,stderr=subprocess.STDOUT,**isolation)
        with patch.object(engine,'_require_windows'),patch.object(engine,'_launch_worker',side_effect=launch),patch.object(stream_worker,'emit',side_effect=lambda event,**data:events.append(dict(event=event,**data))):
            rc=stream_worker.execute(folder)
        return folder,rc,events

    def test_worker_uses_independent_engine_once_and_integrates_with_history_and_inspector(self):
        folder,rc,events=self.run_fake_engine()
        self.assertEqual(rc,0);self.assertEqual(json.loads((folder/'stream-result.json').read_text())['status'],'SUCCESS')
        self.assertIn('NOTE: completed café',(folder/'run.log').read_text(encoding='utf-8'))
        self.assertIn('1. extract.sas · Whole file',(folder/'run.log').read_text(encoding='utf-8'))
        observed=json.loads((folder/'execution/observed.json').read_text(encoding='utf-8'))
        self.assertEqual(observed['text'],(folder/'source.sas').read_text(encoding='utf-8'))
        self.assertEqual(len([e for e in events if e['event']=='start']),1)
        history=self.app.history();self.assertEqual(len(history),1);self.assertEqual(history[0]['status'],'SUCCESS')
        details=self.app.details(self.app.relative(folder));self.assertTrue(details['whole_run']);self.assertEqual(details['status'],'SUCCESS')
        self.assertTrue(any(f['name']=='stream.json' for f in details['files']))
        self.assertTrue(any(f['is_log'] and f['name'].endswith('sas.log') for f in details['files']))
        outputs=[f for f in details['files'] if not f.get('internal') and not f.get('is_log')]
        self.assertTrue(outputs[0]['name'].replace('\\','/').startswith('execution/html/'))
        self.assertTrue(next(f for f in details['files'] if f['name'].replace('\\','/')=='execution/worker.txt')['is_log'])
        self.assertEqual(next(f for f in details['files'] if f['name']=='run.log')['is_log'],True)

    def test_stream_errors_and_partial_exports_never_become_success(self):
        self.doc['export_tables']=True
        folder,rc,_=self.run_fake_engine({'failed_table':0})
        self.assertEqual(rc,1);self.assertEqual(self.app.history()[0]['status'],'EXPORT_ERROR')
        self.assertIn('Access denied',(folder/'run.log').read_text(encoding='utf-8'))
        self.assertEqual(len(list((folder/'execution/tables').glob('*.xlsx'))),1)

    def test_cancelled_before_execution_does_not_call_sas_and_is_inspectable(self):
        folder,_,_=stream.prepare(self.app,self.doc);(folder/'_cancel.request').touch()
        with patch.object(stream_worker,'emit'),patch.object(stream_worker,'run_sas') as run:
            rc=stream_worker.execute(folder,runner=run)
        self.assertEqual(rc,130);run.assert_not_called();self.assertEqual(self.app.history()[0]['status'],'CANCELLED')

    def test_engine_cancel_file_stops_real_owned_process_and_releases_lock(self):
        folder,_,_=stream.prepare(self.app,self.doc)
        self.egp.write_text(json.dumps({'delay':60}))
        cancel=self.root/'stop.request'
        def launch(request,output):
            isolation={'creationflags':subprocess.CREATE_NEW_PROCESS_GROUP} if os.name=='nt' else {'start_new_session':True}
            return subprocess.Popen([sys.executable,str(ROOT/'tests/fixtures/standalone_engine_worker.py'),str(request)],stdout=output,stderr=subprocess.STDOUT,**isolation)
        timer=threading.Timer(.7,lambda:cancel.touch());timer.start()
        try:
            with patch.object(engine,'_require_windows'),patch.object(engine,'_launch_worker',side_effect=launch):
                result=engine.run_sas(self.egp,folder/'source.sas',folder/'execution',cancel_file=cancel,timeout=15)
        finally:timer.cancel()
        self.assertEqual(result.status,'CANCELLED');self.assertLess(result.elapsed_seconds,10)
        self.assertFalse((folder/'execution/.engine.lock').exists())

    def test_engine_preexisting_cancellation_never_launches_sas(self):
        marker=self.root/'cancel';marker.touch()
        with patch.object(engine,'_require_windows'),patch.object(engine,'_launch_worker') as launch:
            result=engine.run_sas(self.egp,self.source,self.root/'cancelled',cancel_file=marker)
        launch.assert_not_called();self.assertEqual(result.status,'CANCELLED')

    def test_failed_worker_launch_is_saved_and_cannot_leave_a_running_task(self):
        facade=SimpleNamespace(name='nt',environ=os.environ)
        with patch.object(ui,'os',facade),patch.object(ui.subprocess,'CREATE_NEW_PROCESS_GROUP',512,create=True),patch.object(ui.subprocess,'CREATE_NO_WINDOW',0x08000000,create=True),patch.object(ui.subprocess,'Popen',side_effect=OSError('Python missing')):
            with self.assertRaisesRegex(OSError,'Python missing'):self.app.launch(dict(action='stream',stream=self.doc))
        self.assertEqual(len(self.app.commands),1);self.assertEqual(next(iter(self.app.commands.values()))['status'],'FAILED')
        self.assertEqual(self.app.history()[0]['status'],'FAILED');self.assertFalse(self.app.uploading);self.assertEqual(self.app.processes,{})

    def test_exact_stream_history_estimate_ignores_names_but_not_code_order_or_export_options(self):
        folder,_,_=self.run_fake_engine()
        record=json.loads((folder/'stream-result.json').read_text());record['elapsed_seconds']=123
        (folder/'stream-result.json').write_text(json.dumps(record))
        self.assertEqual(stream.estimate(self.app,{**self.doc,'name':'Renamed'})['seconds'],123)
        self.assertIsNone(stream.estimate(self.app,{**self.doc,'export_html':False})['seconds'])
        # A whole stream duration must never be reused as a one-file timing.
        self.assertEqual(duration_history(self.root,[],self.app.run_bases(),self.app.resolve),[])

    def test_partial_estimates_sum_sequential_work_and_expose_unknown_parts(self):
        second=self.row('next');second.update(mode='lines',row_start=3,row_end=3)
        with patch.object(self.app,'timing_samples',return_value=[{'key':('extract.sas','',0,0),'seconds':50,'time':10}]):
            result=stream.estimate(self.app,{**self.doc,'rows':[self.row('first'),second]})
        self.assertEqual(result['seconds'],50);self.assertEqual(result['unknown'],1)
        remaining=remaining_work([dict(action='stream',status='RUNNING',tasks={'stream':dict(status='RUNNING',started=100,estimate=result)})],[],[],110)
        self.assertEqual(remaining['seconds'],40);self.assertEqual(remaining['unknown'],1)

    def test_combine_workbooks_keeps_unicode_dates_strings_and_unique_sheet_names(self):
        from openpyxl import Workbook,load_workbook
        from datetime import datetime
        paths=[]
        for index,name in enumerate(('really_long_table_name_abcdefghij','really_long_table_name_abcdefghik')):
            path=self.root/(name+'.xlsx');workbook=Workbook();sheet=workbook.active
            sheet.append(['City','Expression','When']);sheet.append(['Lisboa café',None,datetime(2026,10,21)])
            sheet['B2']='=literal';sheet['B2'].data_type='s'
            workbook.save(path);workbook.close();paths.append(path)
        target=self.root/'combined.xlsx';mapping=stream.combine_tables(paths,target)
        result=load_workbook(target);self.assertEqual(len(result.sheetnames),2);self.assertNotEqual(*result.sheetnames)
        self.assertEqual(result.worksheets[0]['A2'].value,'Lisboa café');self.assertEqual(result.worksheets[0]['B2'].data_type,'s')
        self.assertEqual(result.worksheets[0]['C2'].value,datetime(2026,10,21));self.assertEqual(len(mapping),2);result.close()
        self.assertTrue(all(p.exists() for p in paths))

    def test_merge_failure_preserves_individual_tables_and_is_not_success(self):
        self.doc.update(export_tables=True,consolidate_tables=True)
        with patch.object(stream_worker,'combine_tables',side_effect=ValueError('unreadable worksheet')):
            folder,rc,_=self.run_fake_engine()
        self.assertEqual(rc,1);self.assertEqual(self.app.history()[0]['status'],'EXPORT_ERROR')
        self.assertIn('individual exports retained',(folder/'run.log').read_text(encoding='utf-8'))
        self.assertEqual(len(list((folder/'execution/tables').glob('*.xlsx'))),2)

    def test_stream_active_folder_is_protected_before_first_worker_event(self):
        folder,_,_=stream.prepare(self.app,self.doc)
        self.app.commands['stream']=dict(id='stream',action='stream',status='RUNNING',path=self.app.relative(folder),tasks={},started=time.time())
        self.assertIn(folder,ui_storage.protected_runs(self.app))
        self.assertEqual(self.app.cleanup({'mode':'finished'})['runs'],0)

    def test_launch_selects_stream_worker_and_no_legacy_initialization_or_engine(self):
        # Only this module's OS facade is replaced: pathlib still uses the real host OS.
        facade=SimpleNamespace(name='nt',environ=os.environ)
        with patch.object(ui,'os',facade),patch.object(ui.subprocess,'CREATE_NEW_PROCESS_GROUP',512,create=True),patch.object(ui.subprocess,'CREATE_NO_WINDOW',0x08000000,create=True),patch.object(ui.subprocess,'Popen') as popen,patch.object(ui.threading,'Thread'):
            result=self.app.launch(dict(action='stream',stream=self.doc))
        args,kwargs=popen.call_args
        self.assertEqual(Path(args[0][2]).name,'stream_worker.py');self.assertNotIn(str(self.app.script),args[0])
        self.assertNotIn('PYSAS_SHARED_WORK_PATH',kwargs['env']);self.assertNotIn('PYSAS_UI_CONSOLE',kwargs['env'])
        item=self.app.commands[result['id']];self.assertEqual(item['action'],'stream');self.assertTrue(self.app.resolve(item['path']).is_dir())
        self.app.processes.clear()

    def test_full_backup_roundtrip_contains_stream_sources_and_saved_run(self):
        uploaded=self.app.upload({'target':['stream'],'name':['connection.egp']},io.BytesIO(b'{}'),2)
        folder,_,_=self.run_fake_engine()
        archive,_=self.app.backup()
        try:
            with zipfile.ZipFile(archive) as z:
                self.assertTrue(any(n.startswith('tasks/stream-sources/') for n in z.namelist()))
                self.assertTrue(any(n.endswith('/stream.json') for n in z.namelist()))
            archive.seek(0);raw=archive.read()
        finally:archive.close()
        second=self.root/'restored';second.mkdir();shutil.copy2(ROOT/'pysas.py',second/'pysas.py');other=ui.Workbench(second)
        result=other.backup(io.BytesIO(raw),len(raw));self.assertEqual(result['runs'],1)
        saved=next((other.tasks/'files').glob('*'));restored=stream.import_document(other,(saved/'stream.json').read_bytes())
        self.assertIn('%put café',stream.compose(other,restored)['source'])


@unittest.skipUnless(sys.platform=='win32','Real Windows console/cscript; Enterprise Guide objects substituted')
class StreamWindowsTests(unittest.TestCase):
    def test_stream_composes_once_exports_combines_and_finishes_in_existing_inspector(self):
        from openpyxl import Workbook,load_workbook
        from xml.etree import ElementTree as ET
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);shutil.copy2(ROOT/'pysas.py',root/'pysas.py');app=ui.Workbench(root)
            source=root/'cases.sas';source.write_text('%let ids=10;\n%put &ids.;',encoding='utf-8')
            workbook=Workbook();workbook.active.append(['café',10]);table=root/'table.xlsx';workbook.save(table);workbook.close()
            project=root/'project.egp';settings=ET.Element('project',{'log':'NOTE: stream finished café','expected_server':'SASApp','workbook':str(table)})
            ET.SubElement(settings,'program',name='Connection',server='SASApp');ET.ElementTree(settings).write(project,encoding='utf-8')
            doc=dict(format=stream.FORMAT,name='Windows stream',egp=stream.describe(app,project),server_program='',
                     rows=[dict(id='setup',file=stream.describe(app,source),mode='lines',row_start=None,row_end=1,section=''),
                           dict(id='program',file=stream.describe(app,source),mode='lines',row_start=2,row_end=None,section='')],
                     export_html=True,export_tables=True,consolidate_tables=True)
            folder,_,_=stream.prepare(app,doc)
            bridge=engine._VBS.replace('Option Explicit','Option Explicit\n'+(ROOT/'tests/fixtures/standalone_engine_com.vbs').read_text(encoding='utf-8'),1)
            bridge=bridge.replace('CreateObject("SASEGObjectModel.Application." & version)','New ApplicationStub').replace('In project.CodeCollection\n','In project.CodeCollection.AllItems\n')
            item=dict(id='stream-win',action='stream',name=doc['name'],display_name=doc['name'],status='RUNNING',started=time.time(),path=app.relative(folder),tasks={},args=[])
            app.commands[item['id']]=item
            def observe(event,**values):app.event(item,dict(event=event,time=time.time(),**values))
            with patch.object(engine,'_VBS',bridge),patch.object(stream_worker,'emit',side_effect=observe):
                rc=stream_worker.execute(folder)
            self.assertEqual(rc,0,(folder/'run.log').read_text(encoding='utf-8')+'\n'+ui.read_text(folder/'execution/console.txt'))
            app.complete_worker(item['id'],SimpleNamespace(wait=lambda:rc,stdin=None))
            self.assertEqual(item['status'],'SUCCESS');self.assertEqual(len(item['tasks']),1)
            self.assertEqual(engine._decode((folder/'execution/executed.sas').read_bytes()),(folder/'source.sas').read_text(encoding='utf-8'))
            result=load_workbook(folder/'tables.xlsx');self.assertEqual(len(result.sheetnames),2);self.assertEqual(result.worksheets[0]['A1'].value,'café');result.close()
            detail=app.command_details(item['id']);self.assertEqual(detail['status'],'SUCCESS');self.assertTrue(detail['whole_run'])
            self.assertTrue(any(f['name']=='tables.xlsx' for f in detail['files']))


if __name__ == '__main__':unittest.main()
