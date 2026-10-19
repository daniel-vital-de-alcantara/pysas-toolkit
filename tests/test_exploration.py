import csv
import io
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import Mock,patch
import ui_explore as explore
import ui_servers as servers
import pysas_ui as ui

TOKEN='123456abcdef'

def catalog():
    return dict(libraries=[dict(name='DATA',engines=[],paths=[])],tables=[dict(libname='DATA',name="Cases & ' names",kind='DATA',readable=True)],column_details=[
        dict(libname='DATA',table="Cases & ' names",name='Case ID',type='num',position=1,length=8,format='BEST.'),
        dict(libname='DATA',table="Cases & ' names",name="Text's & value",type='char',position=2,length=512)],refreshes={"columns:DATA.Cases & ' names":dict(captured=1,partial=False)})

def sample_log(query,rows,token=TOKEN,meta=None):
    prefix='PXP'+token+'|';lines=[prefix+'BEGIN']
    for n,row in enumerate(rows,1):
        lines.append(prefix+'ROW|'+str(n))
        for i,value in enumerate(row,1):
            lines.append(prefix+'CELL|'+str(i));raw=str(value).encode('utf-8')
            for start in range(0,len(raw),40):lines.append(prefix+'V|'+raw[start:start+40].ljust(40,b' ').hex().upper())
            lines.append(prefix+'ENDCELL')
        lines.append(prefix+'ENDROW')
    lines.extend([prefix+'META|'+(meta or f'{len(rows)}|0|0|0'),prefix+'DONE'])
    return lines

class ExplorationTests(unittest.TestCase):
    def query(self,**kwargs):return explore.request(catalog(),'DATA',"Cases & ' names",**kwargs)
    def test_typed_filters_and_identifier_escaping(self):
        q=self.query(filters=[dict(column='Case ID',operator='ge',value='100'),dict(column="Text's & value",operator='contains',value="'; %put &secret; /*")],join='OR')
        code=explore.sample_code(TOKEN,q)
        self.assertIn("from DATA.'Cases & '' names'n",code)
        self.assertIn("'Case ID'n >= 100",code)
        self.assertIn(" OR ",code)
        self.assertNotIn('%put &secret',code)
        self.assertIn('outobs=101',code);self.assertIn('work._px_sample',code)
        self.assertNotIn('select *',code)
        for bad in ['1;quit;','NaN','Infinity','1E9999']:
            with self.assertRaises(ValueError):self.query(filters=[dict(column='Case ID',operator='eq',value=bad)])
        for bad in [0,1001,True,'1.5']:
            with self.assertRaises(ValueError):self.query(limit=bad)
        with self.assertRaises(ValueError):self.query(filters=[dict(column='unknown',operator='eq',value='x')])
        with self.assertRaises(ValueError):self.query(filters=[dict(column='Case ID',operator='contains',value='1')])
        with self.assertRaises(ValueError):self.query(join='AND; delete')
    def test_stale_inaccessible_or_incomplete_schema_cannot_be_queried(self):
        for mutate in [lambda c:c['tables'][0].update(readable=False),lambda c:c['column_details'][0].update(stale=True),lambda c:c['refreshes']["columns:DATA.Cases & ' names"].update(partial=True),lambda c:c.update(column_details=[])]:
            c=catalog();mutate(c)
            with self.assertRaises(ValueError):explore.request(c,'DATA',"Cases & ' names")
    def test_unicode_numeric_missing_and_transport_flags(self):
        q=self.query();rows=[['123','café年度'*20],['.A','=formula']]
        result=explore.parse_sample(sample_log(q,rows,meta='2|1|1|1'),TOKEN,q)
        self.assertEqual(result['rows'],[[123.0,rows[0][1]],[None,'=formula']])
        self.assertTrue(result['more_rows']);self.assertTrue(result['transport_limited']);self.assertTrue(result['cells_truncated'])
        for broken in [sample_log(q,rows)[:-1],sample_log(q,rows,meta='99|0|0|0'),sample_log(q,rows)[:-4]]:
            with self.assertRaises(ValueError):explore.parse_sample(broken,TOKEN,q)
        with patch.object(explore,'WIRE_LIMIT',20),self.assertRaises(ValueError):explore.parse_sample(sample_log(q,rows),TOKEN,q)
    def test_column_limit_is_explicit_and_dictionary_missing_metadata_is_visible(self):
        c=catalog();c['column_details']=[dict(c['column_details'][0],name='col'+str(n),position=n) for n in range(105)]
        q=explore.request(c,'DATA',"Cases & ' names")
        self.assertEqual(len(q['columns']),100);self.assertEqual(q['omitted_columns'],5)
        c['tables'].append(dict(libname='DATA',name='Locked',kind='UNKNOWN',readable=False))
        rows=explore.dictionary_rows(c,'DATA','Locked');self.assertEqual(rows[0][8],'Unavailable')
        self.assertEqual(len(explore.dictionary_rows(c,'DATA','col50')),1)
        with self.assertRaises(ValueError):explore.dictionary_rows(c,'OTHER')
    def test_spreadsheet_safe_csv_and_excel_exports(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp).resolve();(root/'pysas.py').touch();app=ui.Workbench(root)
            c=catalog();c['column_details'][0]['label']='=HYPERLINK("x")'
            with patch.object(app,'server_catalog',return_value=c):
                for operation in ['csv','xlsx']:
                    job=app.metadata_jobs.start(dict(snapshot='test',libraries='DATA',operation=operation))
                    deadline=time.monotonic()+5
                    while app.metadata_jobs.get(job['id'])['status']=='WORKING' and time.monotonic()<deadline:time.sleep(.01)
                    done=app.metadata_jobs.get(job['id']);self.assertEqual(done['status'],'READY',done)
                    path=root/done['path']
                    if operation=='csv':
                        with path.open(encoding='utf-8-sig',newline='') as output:rows=list(csv.reader(output))
                        self.assertTrue(rows[1][3].startswith("'="))
                    else:
                        from openpyxl import load_workbook
                        book=load_workbook(path);cell=book.active.cell(2,4)
                        self.assertEqual(cell.data_type,'s');self.assertTrue(cell.value.startswith('='));book.close()
    def test_workbench_saves_preview_result_separately_from_catalog(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp).resolve();(root/'pysas.py').touch();(root/'project.egp').touch();app=ui.Workbench(root)
            app.commands['parent']=dict(id='parent',server=dict(template=str(root/'project.egp'),label='Demo',lib='',use_init=False))
            process=Mock(stdin=None);process.wait.return_value=0
            with patch.object(app,'server_catalog',return_value=catalog()),patch.object(app,'arguments',return_value=('server-refresh',['runner','run'])),patch.object(ui.subprocess,'Popen',return_value=process),patch.object(ui.threading,'Thread'):
                identifier=app.launch(dict(action='server-refresh',level='preview',parent_snapshot='parent',libraries='DATA',table="Cases & ' names"))['id']
            item=app.commands[identifier];folder=root/'runner/runs/sample';(folder/'logs').mkdir(parents=True)
            query=item['server']['query'];(folder/'logs/sample.log').write_text('\n'.join(sample_log(query,[[1,'hello']],item['server']['token'])))
            item['tasks']['t']=dict(status='SUCCESS',path=app.relative(folder));app.complete_worker(identifier,process)
            self.assertEqual(item['status'],'SUCCESS');self.assertNotIn('snapshot',item['server'])
            saved=json.loads((root/item['server']['sample']).read_text());self.assertEqual(saved['rows'],[[1.0,'hello']])
            self.assertTrue((root/item['server']['sample_csv']).exists())

    def test_dictionary_collection_is_serial_and_advances_parent_snapshot(self):
        from types import SimpleNamespace
        c=catalog();c['column_details']=[];c['refreshes']={}
        c['tables'].append(dict(libname='DATA',name='Other',kind='DATA',readable=True))
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp).resolve();(root/'pysas.py').touch();app=ui.Workbench(root);calls=[]
            def launch(data):
                calls.append(data);key='done'+str(len(calls));app.commands[key]=dict(status='SUCCESS',server=dict(snapshot='saved'));return dict(id=key)
            with patch.object(app,'server_catalog',return_value=c),patch.object(app,'launch',side_effect=launch),patch.object(explore,'os',SimpleNamespace(name='nt')):
                job=app.metadata_jobs.start(dict(snapshot='initial',libraries='DATA',operation='collect'))
                deadline=time.monotonic()+3
                while app.metadata_jobs.get(job['id'])['status']=='WORKING' and time.monotonic()<deadline:time.sleep(.01)
            self.assertEqual(app.metadata_jobs.get(job['id'])['done'],2)
            self.assertEqual(calls[0]['parent_snapshot'],'initial');self.assertEqual(calls[1]['parent_snapshot'],'done1')

    def test_metadata_http_reads_and_exports_require_existing_catalog(self):
        import http.client,threading
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp).resolve();(root/'pysas.py').touch();server=ui.make_server(root,0)
            artifact=server.app.storage/'artifacts'/'saved';artifact.mkdir(parents=True)
            ui.atomic_json(artifact/'catalog.json',catalog())
            server.app.commands['saved']=dict(id='saved',started=1,status='SUCCESS',name='Metadata',action='server-refresh',tasks={},server=dict(snapshot=server.app.relative(artifact/'catalog.json')))
            thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
            connection=http.client.HTTPConnection('127.0.0.1',server.server_port,timeout=3)
            try:
                connection.request('GET','/api/dictionary?id=saved');response=connection.getresponse();body=json.loads(response.read());self.assertEqual(response.status,200);self.assertEqual(len(body['rows']),2)
                connection.request('GET','/api/server-preview?id=saved');response=connection.getresponse();self.assertEqual(response.status,400);response.read()
                connection.request('POST','/api/metadata-job',body=json.dumps(dict(snapshot='saved',operation='csv')),headers={'Content-Type':'application/json'});response=connection.getresponse();self.assertEqual(response.status,403);response.read()
            finally:connection.close();server.shutdown();thread.join(3);server.server_close()

    def test_catalog_reads_do_not_invert_merge_and_command_locks(self):
        import threading
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp).resolve();(root/'pysas.py').touch();app=ui.Workbench(root)
            target=app.storage/'artifacts'/'snapshot'/'catalog.json';target.parent.mkdir(parents=True);ui.atomic_json(target,catalog())
            app.commands['saved']=dict(server=dict(snapshot=app.relative(target)))
            finished=threading.Event()
            def read():
                with app.lock:app.server_catalog('saved')
                finished.set()
            with app.catalog_lock:
                thread=threading.Thread(target=read);thread.start()
                completed=finished.wait(1)
            thread.join(2);self.assertTrue(completed,'Catalog reads must not wait for a merge while holding the command lock.')
