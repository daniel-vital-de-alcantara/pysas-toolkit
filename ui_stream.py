"""Ordered source composition and portable Run stream documents.

Only stream_worker invokes engine.py. Validation/export/estimation never run SAS.
"""
from __future__ import annotations

import base64
import binascii
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import re
import statistics
import uuid

from engine import _decode
from pysas import select_graph_text, shared_work_settings
from ui_data import estimate_task

FORMAT = "pysas-run-stream-2"
LEGACY_FORMAT = "pysas-run-stream-1"
INLINE_LIMIT = 1024**2
JSON_LIMIT = 128 * 1024**2
SOURCE_LIMIT = 16 * 1024**2
EGP_LIMIT = 64 * 1024**2
COMPILED_LIMIT = 64 * 1024**2


def label(value, default="", limit=120):
    if not isinstance(value, str) or len(value) > limit or any(ord(c) < 32 for c in value):
        raise ValueError(f"Use a name of up to {limit} characters without control characters.")
    return value.strip() or default


def file_ref(value, suffix):
    if value is None: return None
    if not isinstance(value, dict) or not isinstance(value.get("path"), str) or not value["path"]:
        raise ValueError("Invalid stream file reference.")
    name = label(value.get("name", ""), Path(value["path"]).name, 200)
    if Path(name).name != name or re.search(r'[\\<>:"|?*]', name) or Path(name).suffix.lower() != suffix:
        raise ValueError("Choose a " + suffix + " file.")
    digest = value.get("sha256", "")
    if not isinstance(digest, str) or not re.fullmatch(r"[0-9a-f]{64}", digest):
        raise ValueError("Stream source is missing its checksum. Choose the file again.")
    return dict(path=value["path"], name=name, sha256=digest)


def selection(value, *, legacy=False):
    mode = value.get("mode", "whole")
    if mode not in {"whole", "lines", "section"}: raise ValueError("Choose whole file, lines or sections.")
    bounds = {}
    for key in ("row_start", "row_end"):
        number = value.get(key)
        if number is not None and (type(number) is not int or not 1 <= number <= 100000000):
            raise ValueError("Line numbers must be positive whole numbers or blank.")
        bounds[key] = number
    if mode == "lines" and bounds["row_start"] and bounds["row_end"] and bounds["row_start"] > bounds["row_end"]:
        raise ValueError("The first line is after the last line.")
    section = label(value.get("section", ""), limit=200)
    sections = value.get("sections", [section] if section else [])
    if not isinstance(sections, list) or len(sections) > 1000 or any(not isinstance(item, str) for item in sections):
        raise ValueError("Choose section names from the source.")
    sections = [label(item, limit=200) for item in sections]
    if any(not item for item in sections) or len({item.casefold() for item in sections}) != len(sections):
        raise ValueError("Section choices must be nonempty and unique.")
    if legacy and mode == "lines" and not any(bounds.values()): mode = "whole"
    return dict(mode=mode, section=section, sections=sections, **bounds)


def validate(value):
    if not isinstance(value, dict) or value.get("format") not in {FORMAT, LEGACY_FORMAT}:
        raise ValueError("Choose a Run stream JSON exported by PySAS.")
    legacy = value["format"] == LEGACY_FORMAT
    rows = value.get("rows")
    if not isinstance(rows, list) or len(rows) > 1000: raise ValueError("A stream can contain up to 1,000 cards.")
    result = dict(format=FORMAT, name=label(value.get("name", ""), "Untitled stream"),
                  egp=file_ref(value.get("egp"), ".egp"), server_program=label(value.get("server_program", "")), rows=[],
                  shared_work=shared_work_settings(value.get("shared_work", {})))
    result['connection_mode'] = value.get('connection_mode', 'chosen' if legacy else 'auto')
    if result['connection_mode'] not in {'auto', 'chosen'}: raise ValueError('Choose automatic or explicit connection selection.')
    for key, default in (("export_html", True), ("export_tables", False), ("consolidate_tables", False)):
        setting = value.get(key, default)
        if type(setting) is not bool: raise ValueError(key + " must be a checkbox value.")
        result[key] = setting
    ids, units, inline_bytes = set(), 0, 0
    def identifier(raw):
        if not isinstance(raw, str) or not re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", raw) or raw in ids:
            raise ValueError("Every card and EGP program row needs a unique ID.")
        ids.add(raw); return raw
    for row in rows:
        if not isinstance(row, dict): raise ValueError("Invalid stream card.")
        kind = row.get("source_kind", "file")
        if kind not in {"file", "inline", "egp"}: raise ValueError("Choose a PC file, inline code or an EGP program.")
        code = row.get("code", "")
        if not isinstance(code, str) or len(code.encode("utf-8")) > INLINE_LIMIT:
            raise ValueError("Each inline editor can contain up to 1 MB of SAS code.")
        if kind == "inline": inline_bytes += len(code.encode("utf-8"))
        item = dict(id=identifier(row.get("id")), source_kind=kind, name=label(row.get("name", ""), "Inline code"),
                    file=file_ref(row.get("file"), ".sas"), code=code, project=file_ref(row.get("project"), ".egp"),
                    programs=[], **selection(row, legacy=legacy))
        programs = row.get("programs", [])
        if not isinstance(programs, list) or len(programs) > 1000: raise ValueError("Choose up to 1,000 EGP programs.")
        for program in programs:
            if not isinstance(program, dict) or not isinstance(program.get("program_id"), str) or not 0 < len(program["program_id"]) <= 4096:
                raise ValueError("Invalid EGP program selection.")
            item["programs"].append(dict(id=identifier(program.get("id")), program_id=program["program_id"],
                name=label(program.get("name", ""), "EGP program", 200), flow=label(program.get("flow", ""), "Unassigned", 200),
                file=file_ref(program.get("file"), ".sas"), **selection(program, legacy=legacy)))
        units += len(programs) if kind == "egp" else 1
        result["rows"].append(item)
    if units > 1000: raise ValueError("A stream can submit up to 1,000 files/programs in total.")
    if inline_bytes > SOURCE_LIMIT: raise ValueError("Inline source exceeds 16 MB in total.")
    return result


def scan_source(text):
    lines = text.replace("\r\n", "\n").replace("\r", "\n").splitlines()
    starts, ends = {}, {}
    for index, line in enumerate(lines, 1):
        match = re.fullmatch(r"\s*\* \(please do not delete\) section_(start|end): (.+?);\s*", line, re.I)
        if match:
            key = match[2].casefold()
            (starts if match[1].lower() == "start" else ends).setdefault(key, []).append((index, match[2]))
    sections, ranges, issues = [], {}, []
    for key, occurrences in starts.items():
        name = occurrences[0][1]; closing = ends.get(key, [])
        if len(occurrences) != 1 or len(closing) != 1 or occurrences[0][0] >= closing[0][0]:
            issues.append(name + ": markers are missing, duplicated or out of order.")
        elif not any(line.strip() for line in lines[occurrences[0][0]:closing[0][0]-1]):
            issues.append(name + ": this section has no code.")
        else:
            sections.append(name); ranges[name] = dict(first=occurrences[0][0] + 1, last=closing[0][0] - 1)
    return dict(lines=len(lines), has_code=bool(text.strip()), sections=sections, section_ranges=ranges, section_issues=issues)


def asset_refs(document):
    if document["egp"]: yield document["egp"], ".egp"
    for row in document["rows"]:
        # Inactive card drafts are retained in portable templates too.
        if row["file"]: yield row["file"], ".sas"
        if row["project"]: yield row["project"], ".egp"
        for program in row["programs"]:
            if program["file"]: yield program["file"], ".sas"


def describe(app, path, name=None):
    path = app.input_path(str(path))
    suffix = path.suffix.lower()
    if suffix not in {".sas", ".egp"} or not path.is_file() or path.is_symlink():
        raise ValueError("Choose an existing SAS or EGP file.")
    if path.stat().st_size > (SOURCE_LIMIT if suffix == ".sas" else EGP_LIMIT):
        raise ValueError("Run stream accepts SAS files up to 16 MB and EGP templates up to 64 MB.")
    raw = path.read_bytes()
    result = dict(path=app.relative(path), name=name or path.name, sha256=hashlib.sha256(raw).hexdigest())
    if suffix == ".sas":
        text = _decode(raw)
        result.update(scan_source(text))
    return result


def read_asset(app, ref, suffix):
    path = app.input_path(ref["path"])
    if not path.is_file() or path.suffix.lower() != suffix or path.stat().st_size > (SOURCE_LIMIT if suffix == ".sas" else EGP_LIMIT):
        raise ValueError(ref["name"] + ": source is missing or too large. Choose the file again.")
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != ref["sha256"]:
        raise ValueError(ref["name"] + ": source changed. Choose it again to review the new version.")
    return raw


def assets(app, document):
    result, sas_size = {}, 0
    for ref, suffix in asset_refs(document):
        if ref["sha256"] in result: continue
        raw = read_asset(app, ref, suffix); result[ref["sha256"]] = raw
        if suffix == ".sas": sas_size += len(raw)
    if sas_size > SOURCE_LIMIT: raise ValueError("The stream's distinct SAS files exceed 16 MB in total.")
    if sum(map(len, result.values())) > 80 * 1024**2: raise ValueError("Stream assets exceed 80 MB in total.")
    return result


def shared_work_source(value):
    config = shared_work_settings(value)
    if not config["enabled"]: return ""
    literal = "'" + config["path"].replace("'", "''") + "'"
    return ("/* PySAS shared WORK */\n"
            f"data _null_; if libname('{config['libref']}', {literal}) ne 0 then do; put 'ERROR: PySAS shared WORK is unavailable. Keep its owning session open and check the server path.'; abort abend; end; run;\n"
            f"options user={config['libref']};\n\n")


def source_units(document, raw_assets):
    for index, row in enumerate(document["rows"], 1):
        if row["source_kind"] == "inline":
            yield index, row, row["name"], row["code"], "inline"
        elif row["source_kind"] == "egp":
            if not row["project"]: raise ValueError(f"Card {index}: choose an EGP to extract code from.")
            if not row["programs"]: raise ValueError(f"Card {index}: select at least one EGP program.")
            for program in row["programs"]:
                if not program["file"]: raise ValueError(f"Card {index}: wait for {program['name']} to finish extracting.")
                yield index, program, program["file"]["name"], _decode(raw_assets[program["file"]["sha256"]]), "egp"
        else:
            if not row["file"]: raise ValueError(f"Card {index}: choose a SAS file.")
            yield index, row, row["file"]["name"], _decode(raw_assets[row["file"]["sha256"]]), "file"


def compose(app, value, *, require_egp=False):
    document = validate(value)
    if require_egp and not document["egp"]: raise ValueError("Choose the connection EGP first.")
    if not document["rows"]: raise ValueError("Add at least one source card.")
    raw_assets = assets(app, document)
    chunks, selections, compiled_size, part_hashes = [], [], 0, []
    shared = shared_work_source(document["shared_work"])
    if shared: chunks.append(shared)
    for index, unit, name, text, origin in source_units(document, raw_assets):
        if not text.strip(): raise ValueError(f"Card {index}: enter some SAS code.")
        if unit["mode"] == "lines" and not (unit["row_start"] or unit["row_end"]):
            raise ValueError(f"Card {index}: enter a first or last line, or choose Whole file.")
        sections = unit["sections"] if unit["mode"] == "section" else [""]
        if not sections: raise ValueError(f"Card {index}: select at least one section.")
        for section in sections:
            selection = dict(name=name, section=section,
                             row_start=unit["row_start"] if unit["mode"] == "lines" else None,
                             row_end=unit["row_end"] if unit["mode"] == "lines" else None)
            try: selected = select_graph_text(text, selection)
            except ValueError as exc: raise ValueError(f"Card {index}: {exc}") from exc
            if not selected.strip(): raise ValueError(f"Card {index}: the selected code is empty.")
            scope = "Section " + section if section else f"Lines {unit['row_start'] or 1}–{unit['row_end'] or 'end'}" if unit["mode"] == "lines" else "Whole file"
            caption = (name + " · " + scope).replace("*/", "* /")
            position = len(selections) + 1
            chunk = f"/* PySAS stream row {position}: {caption} */\n{selected}\n\n"
            compiled_size += len(chunk.encode("utf-8"))
            if compiled_size > COMPILED_LIMIT: raise ValueError("Selected code exceeds 64 MB. Split it into smaller streams.")
            digest = hashlib.sha256(selected.encode("utf-8")).hexdigest(); part_hashes.append(digest)
            chunks.append(chunk)
            selections.append(dict(id=unit["id"], card=index, position=position, **selection, scope=scope, origin=origin, sha256=digest))
    fingerprint = hashlib.sha256(json.dumps(dict(parts=part_hashes, shared=shared, egp=document["egp"]["sha256"] if document["egp"] else None,
        server=document["server_program"], html=document["export_html"], tables=document["export_tables"], consolidate=document["consolidate_tables"]), sort_keys=True).encode()).hexdigest()
    return dict(document=document, source="".join(chunks), selections=selections, fingerprint=fingerprint, assets=raw_assets)


def portable(document, raw_assets):
    return {**document, "assets": {digest: base64.b64encode(raw).decode("ascii") for digest, raw in raw_assets.items()}}


def artifact(app, value, kind):
    if kind == "source":
        plan = compose(app, value)
        raw = plan["source"].encode("utf-8"); extension = ".sas"; document = plan["document"]
    else:
        document = validate(value)
        raw = json.dumps(portable(document, assets(app, document)), ensure_ascii=False, indent=2).encode("utf-8")
        extension = ".json"
        if len(raw) > JSON_LIMIT: raise ValueError("Portable stream JSON exceeds 128 MB.")
    with app.lock:
        if app.storage_busy: raise ValueError("Wait for the saved-data operation to finish.")
        folder = app.storage / "artifacts" / uuid.uuid4().hex
        folder.mkdir(parents=True)
        stem = re.sub(r"[^a-zA-Z0-9_-]", "_", document["name"])[:80] or "stream"
        path = folder / (stem + extension)
        path.write_bytes(raw)
    return dict(path=app.relative(path), bytes=len(raw))


def import_document(app, raw):
    if len(raw) > JSON_LIMIT: raise ValueError("Run stream JSON exceeds 128 MB.")
    value = json.loads(raw.decode("utf-8-sig"))
    document = validate(value)
    encoded = value.get("assets")
    if not isinstance(encoded, dict): raise ValueError("Portable stream JSON is missing its embedded files.")
    decoded, total_sas, total_assets = {}, 0, 0
    for ref, suffix in asset_refs(document):
        if ref is None or ref["sha256"] in decoded: continue
        content = encoded.get(ref["sha256"])
        if not isinstance(content, str): raise ValueError(ref["name"] + ": embedded file is missing.")
        try: payload = base64.b64decode(content, validate=True)
        except (ValueError, binascii.Error) as exc: raise ValueError("Invalid embedded stream file.") from exc
        if len(payload) > (SOURCE_LIMIT if suffix == ".sas" else EGP_LIMIT) or hashlib.sha256(payload).hexdigest() != ref["sha256"]:
            raise ValueError(ref["name"] + ": embedded file is too large or its checksum is invalid.")
        if suffix == ".sas":
            _decode(payload); total_sas += len(payload)
        total_assets += len(payload)
        if total_assets > 80 * 1024**2: raise ValueError("Stream assets exceed 80 MB in total.")
        decoded[ref["sha256"]] = (payload, suffix)
    if total_sas > SOURCE_LIMIT: raise ValueError("The stream's distinct SAS files exceed 16 MB in total.")
    # No output is created until the entire document has passed validation.
    with app.upload_lock:
        with app.lock:
            if app.storage_busy: raise ValueError("Wait for the saved-data operation to finish.")
            app.uploading = True
        paths, written = {}, []
        try:
            folder = app.tasks / "stream-sources"; folder.mkdir(exist_ok=True)
            for digest, (payload, suffix) in decoded.items():
                path = folder / (uuid.uuid4().hex + suffix); path.write_bytes(payload); written.append(path)
                paths[digest] = app.relative(path)
            for ref, suffix in asset_refs(document):
                if ref:
                    ref["path"] = paths[ref["sha256"]]
                    if Path(ref["name"]).suffix.lower() == ".sas":
                        ref.update(describe(app, ref["path"], ref["name"]))
        except BaseException:
            for path in written: path.unlink(missing_ok=True)
            raise
        finally: app.uploading = False
    for row in document["rows"]:
        if row["source_kind"] == "inline": row["analysis"] = scan_source(row["code"])
    return document


def estimate(app, value, plan=None):
    plan = plan or compose(app, value)
    matches = []
    for parent, _ in app.run_bases():
        if not parent.is_dir(): continue
        for folder in parent.iterdir():
            record = folder / "stream-result.json"
            if folder.is_symlink() or not record.is_file(): continue
            try:
                data = json.loads(record.read_text(encoding="utf-8"))
                seconds = data.get("elapsed_seconds")
                if data.get("status") == "SUCCESS" and data.get("fingerprint") == plan["fingerprint"] and type(seconds) in (int, float) and math.isfinite(seconds) and seconds > 0:
                    matches.append((record.stat().st_mtime, seconds))
            except (OSError, ValueError): continue
    if matches:
        values = [seconds for _, seconds in sorted(matches, reverse=True)[:10]]
        return dict(seconds=statistics.median(values), samples=len(values), unknown=0, basis="Same selected code, order, EGP and output options; recent successful streams.")
    samples = app.timing_samples()
    rows = [estimate_task(s, samples) for s in plan["selections"]]
    known = [row["seconds"] for row in rows if row["seconds"] is not None]
    unknown = len(rows) - len(known)
    return dict(seconds=sum(known) if known else None, samples=sum(row["samples"] for row in rows), unknown=unknown,
                basis="Sum of comparable program/section/line timings from previous runs. A + is a partial estimate; a single SAS session and data changes can affect duration.")


def prepare(app, value):
    plan = compose(app, value, require_egp=True)
    if plan["document"]["export_tables"] and plan["document"]["consolidate_tables"]:
        import importlib.util
        if importlib.util.find_spec("openpyxl") is None: raise ValueError("Install openpyxl to combine tables into one workbook.")
    forecast = estimate(app, value, plan)
    folder = app.tasks / "files" / (datetime.now().strftime("%Y%m%d_%H%M%S") + "__stream_" + uuid.uuid4().hex[:8])
    folder.mkdir(parents=True)
    try:
        (folder / "source.sas").write_text(plan["source"], encoding="utf-8", newline="")
        doc = json.loads(json.dumps(plan["document"]))
        source_dir = folder / "sources"; source_dir.mkdir()
        for index, (ref, suffix) in enumerate(asset_refs(doc), 1):
            path = source_dir / (str(index).zfill(3) + "__" + ref["name"])
            path.write_bytes(plan["assets"][ref["sha256"]])
            ref["path"] = str(path.relative_to(folder)).replace("\\", "/")
        saved = json.dumps(portable(doc, plan["assets"]), ensure_ascii=False, indent=2)
        if len(saved.encode("utf-8")) > JSON_LIMIT: raise ValueError("Portable stream JSON exceeds 128 MB.")
        (folder / "stream.json").write_text(saved, encoding="utf-8")
        manifest = dict(fingerprint=plan["fingerprint"], selections=plan["selections"], estimate=forecast, document=doc)
        (folder / "stream-plan.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    except BaseException:
        import shutil
        shutil.rmtree(folder)
        raise
    return folder, plan["document"], forecast


def combine_tables(paths, destination):
    """Stream workbook cells to one file; retain original exports on any failure."""
    from openpyxl import load_workbook, Workbook
    from openpyxl.cell import WriteOnlyCell
    workbook = Workbook(write_only=True)
    used, sheets = set(), []
    temporary = destination.with_suffix(".tmp.xlsx")
    try:
        for path in paths:
            source = load_workbook(path, read_only=True, data_only=False)
            try:
                for sheet in source.worksheets:
                    stem = re.sub(r"[\\/*?:\[\]]", "_", path.stem + ("_" + sheet.title if len(source.worksheets) > 1 else ""))[:31] or "Table"
                    name, n = stem, 1
                    while name.casefold() in used:
                        n += 1; name = stem[:31-len(str(n))-1] + "_" + str(n)
                    used.add(name.casefold()); target = workbook.create_sheet(name); sheets.append(dict(file=path.name, sheet=sheet.title, combined_sheet=name))
                    for row in sheet.iter_rows():
                        cells = []
                        for cell in row:
                            copied = WriteOnlyCell(target, value=cell.value)
                            if cell.data_type == "s": copied.data_type = "s"
                            copied.number_format = cell.number_format; cells.append(copied)
                        target.append(cells)
            finally: source.close()
        if not sheets: return []
        workbook.save(temporary); temporary.replace(destination)
        return sheets
    finally:
        workbook.close(); temporary.unlink(missing_ok=True)


def inspector_files(files):
    """Prioritize requested outputs/source; keep bridge diagnostics in the run ZIP."""
    internal = {'status.txt', 'stream-plan.json', 'stream-result.json', 'execution/request.json',
                'execution/result.json', 'execution/_bridge.vbs', 'execution/_automation.xml',
                'execution/source.sas', 'execution/submitted.sas', 'tables-index.json'}
    def rank(file):
        name = file['name'].replace('\\', '/')
        if name == 'execution/worker.txt':
            file.update(is_log=True, log_label='Engine worker')
        if name in internal: file['internal'] = True
        if file.get('is_log'):
            return (0, {'run.log': 0, 'execution/sas.log': 1, 'execution/console.txt': 2, 'execution/worker.txt': 3}.get(name, 4), name)
        return (1, 0 if name == 'tables.xlsx' else 1 if name.startswith('execution/html/') else 2 if name.startswith('execution/tables/') else 3 if name == 'source.sas' else 4 if name == 'stream.json' else 5, name)
    return sorted(files, key=rank)
